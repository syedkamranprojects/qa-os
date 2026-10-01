"""Parse the S&D data dictionary workbook into a machine-readable catalog.

Outputs (in ../knowledge/):
  snd_tables.json          table -> module, description, columns (PK/FK/refs), source sheet
  snd_screen_field_map.json screen caption -> table.column (from the *Mapping sheets)
  snd_tables_index.md      human/grep-friendly index: module -> table -> description
  dd_quality_report.md     inconsistencies found in the source workbook
"""
import openpyxl, glob, json, re, sys, os, collections
sys.stdout.reconfigure(encoding='utf-8')

# Outputs go to ../knowledge. The source workbook is looked up there first, then in its original location.
DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'knowledge')
SRC = [p for p in glob.glob(os.path.join(DIR, 'DD-S&D*.xlsx'))
       + glob.glob(r'C:\MyWork\gias-qa-workspace\regress-master\docs\snd-tables\DD-S&D*.xlsx')
       if not os.path.basename(p).startswith('~$')][0]
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

TABLE_RE = re.compile(r'\b([A-Za-z]{2,5}_[A-Za-z]{2}_[A-Za-z0-9]{2,4}_[A-Za-z0-9_]+|FACT_[A-Z0-9_]+|FN_[A-Z0-9_]+|[A-Z]+_MASTER)\b')
def tname(v):
    if not isinstance(v, str): return None
    m = TABLE_RE.search(v.strip())
    return m.group(1).upper() if m else None
def s(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None and str(v).strip() else None
def yn(v):
    return (s(v) or '').lower().startswith('y')

issues = collections.defaultdict(list)

# ---------- 1. Grouping sheet: module + description per table ----------
grouping = {}
module = None
ws = wb['S&D-EN - Tables Grouping VER 12']
for i, row in enumerate(ws.iter_rows(values_only=True), 1):
    row = list(row) + [None] * 5
    serial, tid, desc, ml = row[1], row[2], row[3], row[4]
    if isinstance(serial, str) and tid is None and serial != 'Serial #' and 'Grouping' not in serial:
        module = s(serial).replace(' - Table List', '').replace('Table List - ', '').replace(' VER 12', '')
        continue
    t = tname(tid)
    if not t:
        continue
    if serial is None:
        issues['grouping'].append(f'row {i}: `{t}` listed without a serial number')
    entry = {'module': module, 'description': s(desc), 'multi_language_table': tname(ml)}
    if t in grouping:
        if grouping[t]['description'] != entry['description']:
            issues['grouping'].append(f'row {i}: `{t}` listed more than once with different descriptions '
                                      f'("{grouping[t]["description"]}" vs "{entry["description"]}")')
        else:
            issues['grouping'].append(f'row {i}: `{t}` listed more than once')
        continue
    grouping[t] = entry

# ---------- 2. Column-level DD sheets ----------
DD_SHEETS = [n for n in wb.sheetnames if n not in (
    'S&D-EN - Tables Grouping VER 12', 'Sheet1', 'Sheet2',
    'Distributor Mapping', 'Outlet-Mapping', 'DSR Mapping', 'Product Mapping')]
# 'DD-S&D-EN-Transactions VER 1' is an older copy of 'DD-EN-Transactions VER 12'; parse it last so VER 12 wins.
DD_SHEETS.sort(key=lambda n: n == 'DD-S&D-EN-Transactions VER 1')

tables = {}
for sheet in DD_SHEETS:
    ws = wb[sheet]
    hdr = None
    cur = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        row = list(row) + [None] * 10
        if hdr is None:
            if any(isinstance(v, str) and v.strip() == 'Column Name' for v in row):
                hdr = {s(v): k for k, v in enumerate(row) if s(v)}
                ci = {k: hdr.get(k) for k in ('Serial #', 'Table Name', 'Column Name', 'Column Description',
                                                'Data Type', 'IS PK', 'IS FK')}
                ci['ref'] = next((hdr[k] for k in hdr if k.startswith('Referenc')), ci['IS FK'] + 1)
            continue
        col = s(row[ci['Column Name']])
        tcell = tname(row[ci['Table Name']])
        serial = row[ci['Serial #']]
        # A serial number marks the start of a table; continuation rows inherit it.
        if serial not in (None, '') and tcell:
            cur = tcell
            if cur in tables and tables[cur]['source_sheet'] != sheet:
                if sheet == 'DD-S&D-EN-Transactions VER 1':
                    cur = None  # older duplicate of a VER 12 table: skip
                    continue
                issues['duplicates'].append(f'`{cur}` defined in both "{tables[cur]["source_sheet"]}" and "{sheet}" (kept first)')
                cur = None
                continue
            tables.setdefault(cur, {'source_sheet': sheet, 'columns': []})
        elif tcell and cur and tcell != cur and serial in (None, ''):
            # Table-name cell on a continuation row disagrees with the table in progress.
            issues['continuation'].append((sheet, i, cur, tcell, col))
        if not cur or not col or not re.match(r'^[A-Za-z0-9_]+$', col):
            continue
        tables[cur]['columns'].append({
            'name': col.upper(),
            'description': s(row[ci['Column Description']]),
            'type': s(row[ci['Data Type']]),
            'pk': yn(row[ci['IS PK']]),
            'fk': yn(row[ci['IS FK']]),
            'references': tname(row[ci['ref']]),
        })

# ---------- 3. Merge ----------
catalog = {}
for t in sorted(set(grouping) | set(tables)):
    g = grouping.get(t, {})
    d = tables.get(t, {})
    catalog[t] = {
        'module': g.get('module') or (d.get('source_sheet') and f'(not in grouping) {d["source_sheet"]}'),
        'description': g.get('description'),
        'multi_language_table': g.get('multi_language_table'),
        'source_sheet': d.get('source_sheet'),
        'primary_key': [c['name'] for c in d.get('columns', []) if c['pk']],
        'columns': d.get('columns', []),
    }
no_cols = [t for t in grouping if t not in tables]
not_grouped = [t for t in tables if t not in grouping]

# Merge check: a table whose sheet block lacked a serial number would have its columns folded into
# the previous table, which shows up as two large groups of non-FK columns with different prefixes.
AUDIT = {'CREATED_BY', 'CREATE_DATE', 'CREATED_DATE', 'MODIFIED_BY', 'MODIFIED_DATE', 'MODIFY_BY', 'MODIFY_DATE'}
prefix_odd = []
for t, e in catalog.items():
    c = collections.Counter(x['name'].split('_')[0] for x in e['columns'] if not x['fk'] and x['name'] not in AUDIT)
    big = sorted((k, v) for k, v in c.items() if v >= 3)
    if len(big) >= 2:
        prefix_odd.append(f'{t}` — ' + ', '.join(f'{k}_* ×{v}' for k, v in big) + ' `')

# ---------- 4. Screen caption -> column maps ----------
screen_map = {}
for sheet in ('Distributor Mapping', 'Outlet-Mapping', 'DSR Mapping', 'Product Mapping'):
    ws = wb[sheet]
    hdr = None
    entries = []
    title = None
    for row in ws.iter_rows(values_only=True):
        row = list(row) + [None] * 8
        if title is None:
            title = next((s(v) for v in row if s(v)), None)
        if hdr is None:
            if any(s(v) == 'Caption/Attribute' for v in row):
                hdr = {s(v): k for k, v in enumerate(row) if s(v)}
            continue
        cap = s(row[hdr['Caption/Attribute']])
        tbl = tname(row[hdr['Table Name']])
        colname = s(row[hdr['Column Name']])
        if not (cap and tbl and colname): continue
        entries.append({'screen_tab': s(row[hdr['Screen Name']]), 'caption': cap,
                        'table': tbl, 'column': colname.upper(), 'type': s(row[hdr['Data Type']])})
    screen_map[title or sheet] = entries

# ---------- 5. Write outputs ----------
meta = {'source': os.path.basename(SRC), 'tables': len(catalog),
        'tables_with_columns': len(tables), 'columns': sum(len(v['columns']) for v in tables.values())}
with open(os.path.join(DIR, 'snd_tables.json'), 'w', encoding='utf-8') as f:
    json.dump({'_meta': meta, 'tables': catalog}, f, indent=1, ensure_ascii=False)
with open(os.path.join(DIR, 'snd_screen_field_map.json'), 'w', encoding='utf-8') as f:
    json.dump({'_meta': {'source': os.path.basename(SRC)}, 'screens': screen_map}, f, indent=1, ensure_ascii=False)

by_mod = collections.defaultdict(list)
for t, e in catalog.items():
    by_mod[e['module'] or '(unknown)'].append(t)
with open(os.path.join(DIR, 'snd_tables_index.md'), 'w', encoding='utf-8') as f:
    f.write(f'# S&D table index\n\nGenerated from `{meta["source"]}` by `dd_build.py` — do not edit by hand.\n'
            f'{meta["tables"]} tables, {meta["columns"]} columns. Full column detail: `snd_tables.json`.\n')
    for m in sorted(by_mod, key=lambda k: (k.startswith('(not'), k)):
        f.write(f'\n## {m}\n\n| Table | Cols | Description |\n|---|---|---|\n')
        for t in sorted(by_mod[m]):
            e = catalog[t]
            f.write(f'| `{t}` | {len(e["columns"]) or "—"} | {e["description"] or ""} |\n')

cont_by_sheet = collections.Counter(x[0] for x in issues['continuation'])
with open(os.path.join(DIR, 'dd_quality_report.md'), 'w', encoding='utf-8') as f:
    f.write(f'# Data dictionary quality report\n\nSource: `{meta["source"]}`\n\n')
    f.write(f'## Listed in grouping sheet but no column definitions ({len(no_cols)})\n\n')
    f.write(''.join(f'- `{t}` — {grouping[t]["module"]}\n' for t in sorted(no_cols)) or '- none\n')
    f.write(f'\n## Column definitions but not in grouping sheet ({len(not_grouped)})\n\n')
    f.write(''.join(f'- `{t}` — sheet "{tables[t]["source_sheet"]}"\n' for t in sorted(not_grouped)) or '- none\n')
    f.write('\n## Continuation rows whose Table Name disagrees with the table being defined\n\n'
            'The parser assigns these columns to the table that started at the last serial number '
            '(the Table Name cell looks copy-pasted). Count by sheet:\n\n')
    f.write(''.join(f'- {k}: {v}\n' for k, v in cont_by_sheet.most_common()) or '- none\n')
    f.write('\nFirst 25 examples (sheet, row, assigned table, cell says, column):\n\n')
    f.write(''.join(f'- {a} r{b}: `{c}` ← cell `{d}` ({e})\n' for a, b, c, d, e in issues['continuation'][:25]))
    f.write(f'\n## Tables with two large groups of column prefixes ({len(prefix_odd)})\n\n'
            'Usually legitimate (e.g. `PANC_*` ancestor columns not flagged as FK), but it could also mean '
            'a following table had no serial number and was merged in. Verify against the live DB.\n\n')
    f.write(''.join(f'- `{t}\n' for t in sorted(prefix_odd)).replace(' `\n', '\n') or '- none\n')
    for k in ('grouping', 'duplicates'):
        f.write(f'\n## {k.title()} issues ({len(issues[k])})\n\n' + (''.join(f'- {x}\n' for x in issues[k]) or '- none\n'))

print(json.dumps(meta))
print('modules:', {m: len(v) for m, v in by_mod.items()})
print('no_cols', len(no_cols), 'not_grouped', len(not_grouped), 'continuation', dict(cont_by_sheet),
      'prefix_odd', len(prefix_odd), 'grouping_issues', len(issues['grouping']), 'dups', len(issues['duplicates']))
