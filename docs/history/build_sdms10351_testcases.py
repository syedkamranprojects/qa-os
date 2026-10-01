from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.colors import Color

OUT = r'C:\MyWork\gias-qa-workspace\regress-master\data\SDMS-10351 Auto Substitute Exclusion Policy - Test Cases.xlsx'
SCREEN = 'Auto Substitute Policy'
VERSION = 'SDMS-10351'
STATUS = 'Not Executed'

PRE = (
    "1) Distributor level user must exist with all assigned roles along with location.\n\n"
    "2) Application URL must be accessible.\n\n"
    "3) Auto Substitute Policy screen must be accessible to the user.\n\n"
    "4) Channel hierarchy (Outlet Subtypes) must exist, be Active and associated with outlets.\n\n"
    "5) Base Auto Substitution must be configured for the SKUs used in the test."
)
def pre(*extra):
    t = PRE
    for i, e in enumerate(extra, 6):
        t += f"\n\n{i}) {e}"
    return t

LOGIN = ("Login to the BackOffice DMS application with Distributor user",
         "Application should be logged in successfully")
NAV = ("Navigate to Auto Substitute Policy",
       "Auto Substitute Policy screen should open successfully with the policy grid")
DL = ("Select the policy from grid and click on Download",
      "Excel file should be downloaded successfully")
def UP(msg):
    return ("Click on Upload, select the edited Excel file and upload", msg)
ORDER_DCODE = ("Create an order in DCODE BackOffice for the outlet with the SKU configured for auto substitution",
               None)
ORDER_MOB = ("Sync an order from Mobility for the outlet with the SKU configured for auto substitution",
             "Order should land in DCODE successfully")

# (section title, [ (description, prerequisite, [(step, expected), ...]) ])
SECTIONS = [
 ("Screen Access and Policy Grid", [
  ("Verify distributor user can navigate to the new Auto Substitute Policy option",
   pre(), [LOGIN, NAV]),
  ("Verify grid displays current policies (Start Date <= current date <= End Date) of the logged-in distributor",
   pre("A current policy must exist for the distributor."),
   [LOGIN, NAV,
    ("Check the policy grid", "Current policy should be displayed in the grid with its Start Date and End Date")]),
  ("Verify grid displays future policies (Start Date > current date) of the logged-in distributor",
   pre("A future policy must exist for the distributor."),
   [LOGIN, NAV,
    ("Check the policy grid", "Future policy should be displayed in the grid")]),
  ("Verify grid does not display expired policies (End Date < current date)",
   pre("An expired policy must exist for the distributor."),
   [LOGIN, NAV,
    ("Check the policy grid", "Expired policy should not be displayed in the grid")]),
  ("Verify grid displays only the policies defined for the logged-in distributor",
   pre("Policies must exist for at least two distributors."),
   [LOGIN, NAV,
    ("Check the policy grid", "Only the logged-in distributor's policies should be displayed; other distributors' policies should not be visible")]),
 ]),
 ("Download of Auto Substitute Policy Excel", [
  ("Verify downloaded Excel contains the required columns",
   pre(), [LOGIN, NAV, DL,
    ("Open the downloaded Excel file",
     "Excel should contain columns: 1) Policy Code, 2) Entity Code, 3) Entity Type, 4) Channel Hierarchy")]),
  ("Verify Policy Code and Entity Code are pre-filled in downloaded Excel",
   pre(), [LOGIN, NAV, DL,
    ("Open the downloaded Excel file and check Policy Code and Entity Code",
     "Policy Code should be pre-filled as per organization setup; Entity Code should be pre-filled with the logged-in DT code")]),
  ("Verify downloaded Excel of a current/future policy contains previously uploaded data",
   pre("Outlet subtypes must already be uploaded against a current or future policy."),
   [LOGIN, NAV, DL,
    ("Open the downloaded Excel file",
     "Previously uploaded rows should be present with Entity Type pre-filled and Channel Hierarchy description of the Outlet Subtype pre-filled")]),
  ("Verify Entity Type is empty in downloaded Excel when nothing was uploaded previously",
   pre("Policy must have no previously uploaded outlet subtypes."),
   [LOGIN, NAV, DL,
    ("Open the downloaded Excel file and check Entity Type",
     "Entity Type should be blank for the user to fill")]),
 ]),
 ("Upload of Auto Substitute Policy (Excel Validations)", [
  ("Verify successful upload with valid Policy Code, Entity Code, Entity Type and active Channel Hierarchy",
   pre("Future policy must exist."),
   [LOGIN, NAV, DL,
    ("Enter valid Entity Type and active Channel Hierarchy (Outlet Subtype) code",
     "Data should be entered successfully"),
    UP("File should be uploaded successfully with success message; outlet subtype should be saved against the policy")]),
  ("Verify upload is rejected when Entity Code is blank",
   pre(), [LOGIN, NAV, DL,
    ("Remove Entity Code from a row", "Entity Code should be cleared"),
    UP("Upload should be rejected with validation message that Entity Code is mandatory")]),
  ("Verify upload is rejected when Entity Type is blank",
   pre(), [LOGIN, NAV, DL,
    ("Leave Entity Type blank for a row", "Entity Type should be blank"),
    UP("Upload should be rejected with validation message that Entity Type is mandatory")]),
  ("Verify Channel Hierarchy is non-mandatory (upload with blank Channel Hierarchy description)",
   pre(), [LOGIN, NAV, DL,
    ("Fill mandatory fields and leave the Channel Hierarchy description blank", "Data should be entered successfully"),
    UP("Upload should not fail due to blank Channel Hierarchy description")]),
  ("Verify upload is rejected when Channel Hierarchy code does not exist",
   pre(), [LOGIN, NAV, DL,
    ("Enter a non-existent Channel Hierarchy code", "Data should be entered"),
    UP("Upload should be rejected with message that Channel Hierarchy code does not exist")]),
  ("Verify upload is rejected when Channel Hierarchy code is In-Active",
   pre("An In-Active outlet subtype must exist."),
   [LOGIN, NAV, DL,
    ("Enter the In-Active Channel Hierarchy code", "Data should be entered"),
    UP("Upload should be rejected with message that Channel Hierarchy code must be Active")]),
  ("Verify upload is rejected when Entity Code is changed to another / invalid DT code",
   pre(), [LOGIN, NAV, DL,
    ("Change Entity Code to a DT code other than the logged-in distributor", "Data should be entered"),
    UP("Upload should be rejected with message for invalid Entity Code")]),
  ("Verify upload is rejected when pre-filled Policy Code is modified",
   pre(), [LOGIN, NAV, DL,
    ("Change the pre-filled Policy Code to an invalid value", "Data should be entered"),
    UP("Upload should be rejected with message for invalid Policy Code")]),
 ]),
 ("Overlapping Policy Validation (DT + Outlet Subtype)", [
  ("Verify upload is rejected when the same outlet subtype already exists in an overlapping policy",
   pre("Policy A with Outlet Subtype X must exist.", "Policy B must exist with dates overlapping Policy A."),
   [LOGIN, NAV,
    ("Select Policy B from grid and click on Download", "Excel file should be downloaded successfully"),
    ("Enter Outlet Subtype X", "Data should be entered"),
    UP("Upload should be rejected with message that overlapping policy dates for the same outlet subtype are not allowed")]),
  ("Verify upload is accepted for a different outlet subtype in an overlapping policy",
   pre("Policy A with Outlet Subtype X must exist.", "Policy B must exist with dates overlapping Policy A."),
   [LOGIN, NAV,
    ("Select Policy B from grid and click on Download", "Excel file should be downloaded successfully"),
    ("Enter Outlet Subtype Y (not in Policy A)", "Data should be entered"),
    UP("File should be uploaded successfully")]),
  ("Verify new policy for the same outlet subtype is accepted after end-dating the existing policy",
   pre("Current Policy A with Outlet Subtype X must exist."),
   [LOGIN, NAV,
    ("Update End Date of Policy A to a date before the Start Date of new Policy B and save", "End Date should be updated successfully"),
    ("Select Policy B from grid and click on Download", "Excel file should be downloaded successfully"),
    ("Enter Outlet Subtype X", "Data should be entered"),
    UP("File should be uploaded successfully as the dates no longer overlap")]),
 ]),
 ("Updation of Auto Substitute Policy", [
  ("Verify future policy can be fully updated",
   pre("A future policy must exist."),
   [LOGIN, NAV,
    ("Select the future policy from grid", "Policy details should be displayed"),
    ("Update Start Date and End Date and save", "Policy should be updated successfully"),
    ("Download, change the outlet subtypes and upload", "Outlet subtypes should be updated successfully")]),
  ("Verify only End Date is editable for a current policy",
   pre("A current policy must exist."),
   [LOGIN, NAV,
    ("Select the current policy from grid", "Policy details should be displayed"),
    ("Check editable fields", "Only End Date should be editable; all other fields should be disabled"),
    ("Update End Date to a valid future date and save", "End Date should be updated successfully")]),
  ("Verify Start Date cannot be changed for a current policy",
   pre("A current policy must exist."),
   [LOGIN, NAV,
    ("Select the current policy from grid", "Policy details should be displayed"),
    ("Try to change Start Date", "Start Date should not be editable / system should not allow the change")]),
  ("Verify outlet subtypes cannot be changed via upload for a current policy",
   pre("A current policy with uploaded outlet subtypes must exist."),
   [LOGIN, NAV,
    ("Select the current policy and click on Download", "Excel file should be downloaded successfully"),
    ("Add / remove an outlet subtype row", "Data should be changed"),
    UP("Upload should be rejected since only End Date of a current policy can be edited")]),
  ("Verify End Date cannot be earlier than Start Date",
   pre("A future policy must exist."),
   [LOGIN, NAV,
    ("Select the policy and set End Date earlier than Start Date, then save", "System should show validation message and not save the policy")]),
 ]),
 ("Auto Substitution Behaviour in DCODE (Acceptance Criteria)", [
  ("Verify auto substitution is NOT applied on order created in DCODE for an outlet of an excluded outlet subtype within policy dates",
   pre("Current policy must exist with Outlet Subtype X for the DT.", "Outlet must belong to Outlet Subtype X."),
   [LOGIN, (ORDER_DCODE[0], "Order should be created successfully"),
    ("Check order lines", "Auto substitution should not be applied; ordered SKU should remain unchanged")]),
  ("Verify auto substitution is NOT applied on order landing from Mobility for an outlet of an excluded outlet subtype within policy dates",
   pre("Current policy must exist with Outlet Subtype X for the DT.", "Outlet must belong to Outlet Subtype X."),
   [ORDER_MOB, LOGIN,
    ("Open the synced order and check order lines", "Auto substitution should not be applied; ordered SKU should remain unchanged")]),
  ("Verify auto substitution IS applied for an outlet whose subtype is not in the policy",
   pre("Current policy must exist with Outlet Subtype X for the DT.", "Outlet must belong to Outlet Subtype Y (not excluded)."),
   [LOGIN, (ORDER_DCODE[0], "Order should be created successfully"),
    ("Check order lines", "Auto substitution should be applied as per the existing logic")]),
  ("Verify auto substitution IS applied when order date is before policy Start Date",
   pre("Future policy must exist with Outlet Subtype X.", "Outlet must belong to Outlet Subtype X."),
   [LOGIN, (ORDER_DCODE[0] + " (order date before policy Start Date)", "Order should be created successfully"),
    ("Check order lines", "Auto substitution should be applied")]),
  ("Verify auto substitution IS applied when order date is after policy End Date",
   pre("Expired policy must exist with Outlet Subtype X.", "Outlet must belong to Outlet Subtype X."),
   [LOGIN, (ORDER_DCODE[0] + " (order date after policy End Date)", "Order should be created successfully"),
    ("Check order lines", "Auto substitution should be applied")]),
  ("Verify boundary: auto substitution is NOT applied on policy Start Date and End Date",
   pre("Policy must exist with Outlet Subtype X.", "Outlet must belong to Outlet Subtype X."),
   [LOGIN, (ORDER_DCODE[0] + " with order date = Start Date", "Order should be created; auto substitution should not be applied"),
    (ORDER_DCODE[0] + " with order date = End Date", "Order should be created; auto substitution should not be applied")]),
  ("Verify exclusion is distributor specific (same outlet subtype under another DT is substituted)",
   pre("Policy with Outlet Subtype X must exist for DT-1 only.", "Outlet of Subtype X must exist under DT-2."),
   [("Login to the BackOffice DMS application with DT-2 user", "Application should be logged in successfully"),
    (ORDER_DCODE[0], "Order should be created successfully"),
    ("Check order lines", "Auto substitution should be applied since the policy belongs to DT-1")]),
  ("Verify auto substitution resumes after End Date of a current policy is shortened",
   pre("Current policy must exist with Outlet Subtype X.", "Outlet must belong to Outlet Subtype X."),
   [LOGIN, NAV,
    ("Update End Date of the policy to current date and save", "End Date should be updated successfully"),
    (ORDER_DCODE[0] + " with order date after the new End Date", "Order should be created successfully"),
    ("Check order lines", "Auto substitution should be applied")]),
  ("Regression: verify previous SKU exclusion list logic is removed",
   pre("SKU must be present in the old SKU exclusion list.", "Outlet subtype must NOT be in any Auto Substitute exclusion policy."),
   [LOGIN, (ORDER_DCODE[0], "Order should be created successfully"),
    ("Check order lines", "Auto substitution should be applied; old SKU exclusion list should no longer block substitution")]),
 ]),
]

# Case-level expected results where the last step's expectation alone is not the full outcome.
EXPECTED = {
  "Verify distributor user can navigate to the new Auto Substitute Policy option":
    "Auto Substitute Policy option should be available in the menu and the screen should open with the policy grid",
  "Verify future policy can be fully updated":
    "Start Date, End Date and outlet subtypes of the future policy should be updated successfully",
  "Verify only End Date is editable for a current policy":
    "Only End Date should be editable (all other fields disabled) and the updated End Date should be saved successfully",
  "Verify boundary: auto substitution is NOT applied on policy Start Date and End Date":
    "Auto substitution should not be applied for orders dated on the policy Start Date or End Date",
  "Verify new policy for the same outlet subtype is accepted after end-dating the existing policy":
    "After End Date of the existing policy is updated, upload of the same outlet subtype in the new policy should succeed as dates no longer overlap",
  "Verify auto substitution resumes after End Date of a current policy is shortened":
    "Auto substitution should be applied for orders dated after the updated End Date",
}

# ---- styles (mirroring the reference template) ----
thin = Side(style='thin')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill('solid', fgColor=Color(theme=4, tint=0.5999938962981048))
SEC_FILL = PatternFill('solid', fgColor='FFFFFF00')
F = Font(name='Calibri', size=11)
FB = Font(name='Calibri', size=11, bold=True)
FT = Font(name='Calibri', size=14, bold=True)
WRAP_C = Alignment(horizontal='center', vertical='center', wrap_text=True)
WRAP_L = Alignment(horizontal='left', vertical='center', wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
for col, w in {'A': 21.3, 'B': 55.4, 'C': 38, 'D': 41.6, 'E': 55.7, 'F': 58.1, 'G': 20, 'H': 33.6, 'I': 27.9}.items():
    ws.column_dimensions[col].width = w

def cell(ref, val, font=F, fill=None, align=WRAP_L, border=True):
    c = ws[ref]
    c.value = val
    c.font = font
    c.alignment = align
    if fill: c.fill = fill
    if border: c.border = BORDER
    return c

cell('A1', 'Test Screen:', FT, border=False, align=Alignment())
cell('A2', 'Screen #', fill=HDR_FILL, align=WRAP_C); cell('B2', 'ScreenName', fill=HDR_FILL)
cell('A3', 1, align=WRAP_C); cell('B3', SCREEN)
cell('A5', 'Test Scenario:', FT, border=False, align=Alignment())
cell('A6', 'Screen#', fill=HDR_FILL, align=WRAP_C); cell('B6', 'Test Scenario Description', fill=HDR_FILL)
cell('A7', SCREEN, align=WRAP_C)
cell('B7', '[PH] Auto Substitute Addendum - Maintain Auto Substitute exclusion list at Distributor + Outlet Subtype level (SDMS-10351)')
ws.row_dimensions[7].height = 69
cell('A9', 'Test Cases:', FT, border=False, align=Alignment())
headers = ['TestCase # ', 'Test Case Description', 'Pre- Requisite', 'Test Steps ', 'Actual Result', 'Expected Result',
           'Test Status (Passed/Failed/Blocked/Not Executed)', 'Test Case Evidence', 'Version']
for i, h in enumerate(headers):
    cell(f'{"ABCDEFGHI"[i]}10', h, FB, HDR_FILL, WRAP_C)
ws.row_dimensions[10].height = 51
ws.freeze_panes = 'A11'

r, tc = 11, 0
for title, cases in SECTIONS:
    ws.merge_cells(f'A{r}:I{r}')
    cell(f'A{r}', title, FB, SEC_FILL, WRAP_C)
    for col in 'BCDEFGHI': ws[f'{col}{r}'].border = BORDER; ws[f'{col}{r}'].fill = SEC_FILL
    r += 1
    for desc, prereq, steps in cases:
        tc += 1
        # Steps are entered manually by QA; keep only the case-level expected outcome.
        expected = EXPECTED.get(desc, steps[-1][1])
        cell(f'A{r}', tc, align=WRAP_C)
        cell(f'B{r}', desc)
        cell(f'C{r}', prereq)
        cell(f'D{r}', None)
        cell(f'E{r}', None)
        cell(f'F{r}', expected)
        cell(f'G{r}', STATUS, align=WRAP_C)
        cell(f'H{r}', None)
        cell(f'I{r}', VERSION, align=WRAP_C)
        ws.row_dimensions[r].height = prereq.count('\n') * 15 + 30
        r += 1

wb.save(OUT)
print('test cases:', tc, 'last row:', r - 1, '->', OUT)
