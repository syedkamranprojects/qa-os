"""Render the team-format test-case workbook from a QA OS run folder.

  python runtime/qaos_workbook.py <run_dir>

Reads requirement.json (optional), cases.json, steps.json (optional).
Writes <run_dir>/<KEY>_TestCases.xlsx:
  Sheet1 — team layout: Test Screen / Test Scenario / Test Cases (TestCase#, Description (ALT- for negatives),
           Pre-Requisite, Test Steps (rendered from steps.json), Actual Result, Expected Result,
           Status (dropdown), Evidence, Version); yellow section bands; frozen header.
  Notes  — ambiguities (question, story quote, default, evidence, answer) and beyond-story cases.
  Automation — per case: readiness, blocked_by, mutates_data, data row.
"""
import json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding='utf-8')
run = sys.argv[1]
load = lambda f: json.load(open(os.path.join(run, f), encoding='utf-8')) if os.path.exists(os.path.join(run, f)) else None
req, cases, steps = load('requirement.json'), load('cases.json'), load('steps.json')

# newest LIVE result per case (exec/<case>/<stamp>/result.json), used for Actual Result / Status / Evidence
live = {}
for case_dir in (os.listdir(os.path.join(run, 'exec')) if os.path.isdir(os.path.join(run, 'exec')) else []):
    stamps = sorted(s for s in os.listdir(os.path.join(run, 'exec', case_dir)) if not s.endswith('-dry'))
    if stamps:
        p = os.path.join(run, 'exec', case_dir, stamps[-1])
        live[case_dir] = (json.load(open(os.path.join(p, 'result.json'), encoding='utf-8')), os.path.relpath(p, run))


def actual(cid):
    if by_case.get(cid, {}).get('readiness') == 'deferred':
        note = 'Deferred: file-upload automation is postponed; run manually or via the agreed alternative.'
        if cid in live:
            res, rel = live[cid]
            note += f" (An earlier automated run is not counted: verdict {res['verdict']}.)"
        return note, 'Not Executed', None
    if cid not in live:
        return None, 'Not Executed', None
    res, rel = live[cid]
    bad = [s for s in res['steps'] if s['status'] != 'pass']
    shots = [f for f in os.listdir(os.path.join(run, rel)) if f.endswith('.png')]
    if res['verdict'] == 'pass':
        text = f"All {len(res['steps'])} steps passed ({res['finished']})."
        status = 'Passed'
    elif res['verdict'] in ('needs-data', 'needs-verb'):
        text, status = f"Blocked at step {bad[0]['n']}: {bad[0]['detail']}", 'Blocked'
    else:
        text, status = f"Step {bad[0]['n']} ({bad[0]['do']}) failed: {bad[0]['detail']}", 'Failed'
    return text, status, '\n'.join(os.path.join(rel, f) for f in shots) or rel
by_case = {c['case']: c for c in (steps or {}).get('cases', [])}
defaults = (steps or {}).get('defaults', {}).get('data', {})

# ---------------------------------------------------------------- step rendering
def val(v, data):
    def sub(m):
        expr = m.group(1).strip()
        ns, _, k = expr.partition('.')
        if ns == 'data' and data.get(k) not in (None, ''):
            return str(data[k])
        if expr.startswith('today'):
            return expr.replace(' ', '').upper().replace('TODAY', 'Today')
        return f'<{k or expr}>'
    return re.sub(r'\{\{(.+?)\}\}', sub, str(v))


def expect_text(e, data):
    e = {k: (val(v, data) if isinstance(v, str) else v) for k, v in e.items()}
    if 'url_contains' in e: return f"Verify the page URL contains {e['url_contains']}"
    if 'breadcrumb' in e: return 'Verify breadcrumb shows ' + ' > '.join(e['breadcrumb'])
    if 'grid' in e and 'columns' in e: return f"Verify {e['grid']} grid shows columns: " + ', '.join(e['columns'])
    if 'grid' in e and 'row' in e: return f"Verify {e['grid']} grid {'contains' if e.get('present', True) else 'does not contain'} row {e['row']}"
    if 'fields_present' in e: return 'Verify fields are shown: ' + ', '.join(e['fields_present'])
    if 'field' in e:
        parts = [f"Verify {e['field']}"] + ([f"is {e['state']}"] if 'state' in e else []) + ([f"= \"{e['value']}\""] if 'value' in e else [])
        return ' '.join(parts)
    if 'button' in e: return f"Verify {e['button']} button is {e['state']}"
    if 'message_contains' in e: return f"Verify message contains \"{e['message_contains']}\""
    return 'Verify ' + json.dumps(e, ensure_ascii=False)


def render(c):
    data = {**defaults, **c.get('data', {})}
    out = []
    for s in c['steps']:
        v = s['do']
        t = {
            'login': lambda: 'Login' + (f" and select distributor {val(s['distributor'], data)}" if s.get('distributor') else ''),
            'logout': lambda: 'Logout',
            'navigate': lambda: f"Navigate to {s['menu']}",
            'enter': lambda: f"Enter {s['field']} = \"{val(s['value'], data)}\"",
            'choose': lambda: f"Choose {s['field']} = \"{val(s['value'], data)}\"",
            'set_date': lambda: f"Set {s['field']} = {val(s['value'], data)}",
            'click': lambda: f"Click {s['button']}",
            'select_row': lambda: f"Select {val(s['value'], data)} in {s['grid']} grid",
            'verify_ui': lambda: expect_text(s['expect'], data),
            'verify_db': lambda: f"Verify in DB ({s['recipe']}): {val(json.dumps(s['expect'], ensure_ascii=False) if not isinstance(s['expect'], str) else s['expect'], data)}",
            'capture': lambda: f"Note the {s['field']} value (as {s['as']})",
            'screenshot': lambda: 'Take screenshot',
            'wait': lambda: f"Wait {s['seconds']}s",
            'download': lambda: f"Click {s['button']} and save the file as {s['as']}",
            'upload': lambda: f"Click {s['button']} and upload {val(s['file'], data)}",
            'edit_excel': lambda: f"In {val(s['file'], data)}: " + '; '.join(
                ', '.join(f"{k} = {val(x, data)}" for k, x in r.items()) for r in s['rows']),
            'manual': lambda: f"[Manual] {s['note']}",
            'capture_message': lambda: 'Note the message shown on screen',
            'harvest_grid': lambda: f"Note the rows listed in the {s.get('grid', 'grid')} grid",
            'harvest_menu': lambda: None, 'harvest_i18n': lambda: None,
        }[v]()
        if t:
            out.append(t)
    return '\n'.join(f'{i}) {t}' for i, t in enumerate(out, 1))

# ---------------------------------------------------------------- workbook
thin = Side(style='thin'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor=Color(theme=4, tint=0.5999938962981048)); SEC = PatternFill('solid', fgColor='FFFFFF00')
F, FB, FT = Font(name='Calibri', size=11), Font(name='Calibri', size=11, bold=True), Font(name='Calibri', size=14, bold=True)
WC = Alignment(horizontal='center', vertical='center', wrap_text=True); WL = Alignment(horizontal='left', vertical='center', wrap_text=True)
TOP = Alignment(horizontal='left', vertical='top', wrap_text=True)

wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
for col, w in {'A': 12, 'B': 48, 'C': 42, 'D': 62, 'E': 40, 'F': 52, 'G': 18, 'H': 34, 'I': 16}.items():
    ws.column_dimensions[col].width = w


def cell(ref, v, font=F, fill=None, align=WL, border=True):
    c = ws[ref]; c.value = v; c.font = font; c.alignment = align
    if fill: c.fill = fill
    if border: c.border = BORDER


key = cases['key']
screen = cases.get('screen_name') or (req or {}).get('title', key)
scenario = cases.get('scenario') or f"{(req or {}).get('title', '')} ({key})"
cell('A1', 'Test Screen:', FT, border=False, align=Alignment())
cell('A2', 'Screen #', fill=HDR, align=WC); cell('B2', 'ScreenName', fill=HDR)
cell('A3', 1, align=WC); cell('B3', screen)
cell('A5', 'Test Scenario:', FT, border=False, align=Alignment())
cell('A6', 'Screen#', fill=HDR, align=WC); cell('B6', 'Test Scenario Description', fill=HDR)
cell('A7', screen, align=WC); cell('B7', scenario); ws.row_dimensions[7].height = 45
cell('A9', 'Test Cases:', FT, border=False, align=Alignment())
heads = ['TestCase # ', 'Test Case Description', 'Pre- Requisite', 'Test Steps ', 'Actual Result', 'Expected Result',
         'Test Status (Passed/Failed/Blocked/Not Executed)', 'Test Case Evidence', 'Version']
for i, h in enumerate(heads):
    cell(f'{"ABCDEFGHI"[i]}10', h, FB, HDR, WC)
ws.row_dimensions[10].height = 45
ws.freeze_panes = 'A11'
dv = DataValidation(type='list', formula1='"Passed,Failed,Blocked,Not Executed"', allow_blank=True)
ws.add_data_validation(dv)

r = 11
for sec in cases['sections']:
    ws.merge_cells(f'A{r}:I{r}'); cell(f'A{r}', sec['title'], FB, SEC, WC)
    for col in 'BCDEFGHI':
        ws[f'{col}{r}'].border = BORDER; ws[f'{col}{r}'].fill = SEC
    r += 1
    for c in sec['cases']:
        title = ('ALT-' if c['kind'] == 'negative' else '') + c['title']
        pre = '\n'.join(f'{i}) {p}' for i, p in enumerate(c['preconditions'], 1))
        st = render(by_case[c['id']]) if c['id'] in by_case else None
        act, status, evid = actual(c['id'])
        for col, v, al in (('A', c['id'], WC), ('B', title, WL), ('C', pre, TOP), ('D', st, TOP), ('E', act, TOP),
                           ('F', c['expected'], TOP), ('G', status, WC), ('H', evid, TOP), ('I', key, WC)):
            cell(f'{col}{r}', v, align=al)
        if status in ('Passed', 'Failed', 'Blocked'):
            ws[f'G{r}'].fill = PatternFill('solid', fgColor={'Passed': 'FFC6EFCE', 'Failed': 'FFFFC7CE', 'Blocked': 'FFFFEB9C'}[status])
        dv.add(f'G{r}')
        lines = max(len(pre.split('\n')) * 1.4, len((st or '').split('\n')) * 1.25, len(c['expected']) / 45 + 1, 3)
        ws.row_dimensions[r].height = min(409, 15 * lines)
        r += 1

# Notes
ns = wb.create_sheet('Notes')
for col, w in {'A': 10, 'B': 60, 'C': 50, 'D': 50, 'E': 50, 'F': 30}.items():
    ns.column_dimensions[col].width = w
row = 1
def nrow(vals, font=F, fill=None):
    global row
    for i, v in enumerate(vals):
        c = ns.cell(row=row, column=i + 1, value=v); c.font = font; c.alignment = TOP; c.border = BORDER
        if fill: c.fill = fill
    row += 1
if req:
    nrow(['Ambiguities in the story — answer column F, or the default is used'], FB); row += 0
    nrow(['Id', 'Question', 'Story quote', 'Default assumption', 'Evidence (live app / DB)', 'Answer'], FB, HDR)
    for a in req.get('ambiguities', []):
        nrow([a['id'], a['question'], a['quote'], a['default'], a.get('evidence', ''), a.get('answer') or ''])
    row += 1
nrow(['Cases testing behaviour the story does not state — BA to confirm or delete'], FB)
nrow(['Case', 'Title', 'Traces'], FB, HDR)
for s in cases['sections']:
    for c in s['cases']:
        if c.get('beyond_story'):
            nrow([c['id'], c['title'], ', '.join(c['traces'])])

# Automation
au = wb.create_sheet('Automation')
for col, w in {'A': 10, 'B': 16, 'C': 60, 'D': 12, 'E': 70}.items():
    au.column_dimensions[col].width = w
for i, h in enumerate(['Case', 'Readiness', 'Blocked by', 'Mutates data', 'Data row'], 1):
    c = au.cell(row=1, column=i, value=h); c.font = FB; c.fill = HDR; c.border = BORDER
for i, (cid, c) in enumerate(by_case.items(), 2):
    for j, v in enumerate([cid, c.get('readiness', ''), '\n'.join(c.get('blocked_by', [])), 'Y' if c.get('mutates_data') else '',
                           json.dumps({**defaults, **c.get('data', {})}, ensure_ascii=False)], 1):
        x = au.cell(row=i, column=j, value=v); x.alignment = TOP; x.border = BORDER
au.freeze_panes = 'A2'

out = os.path.join(run, f'{key}_TestCases.xlsx')
wb.save(out)
n = sum(len(s['cases']) for s in cases['sections'])
print(f'{n} cases, {len(by_case)} with steps -> {out}')
