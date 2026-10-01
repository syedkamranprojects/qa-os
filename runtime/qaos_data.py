"""data-engineer resolver: turn a discovery run's harvested data into ONE data row for the cases of a QA OS run.

  python runtime/qaos_data.py <run_dir> <discovery_result_dir> [--apply]

Reads   <discovery>/data/*.json   (harvest_grid outputs)   and   <discovery>/files/*.xlsx   (downloaded templates)
Writes  <run_dir>/data.json       resolved values with provenance (source file, when)
        --apply: fills only NULL values in steps.json defaults.data (never overwrites a value someone set)
Then re-run the dry run to see which cases are ready.

S&D SDMS-10351 resolver: classifies policies (current/future/expired by Start/End Date vs today), picks active/inactive
channel hierarchy codes, reads Entity Type / Entity Code / column names from the downloaded template.
"""
import datetime as dt, glob, json, os, re, sys

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
run, disc = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
apply = '--apply' in sys.argv
today = dt.date.today()
now = dt.datetime.now().isoformat(timespec='seconds')


def jload(name):
    p = os.path.join(disc, 'data', f'{name}.json')
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else None


def date(v):
    try:
        return dt.datetime.strptime(str(v).strip()[:10], '%Y-%m-%d').date()
    except ValueError:
        return None


def classify(rows):
    out = {'current': [], 'future': [], 'expired': []}
    for r in rows:
        s, e = date(r.get('Start Date')), date(r.get('End Date'))
        if not s or not e:
            continue
        out['future' if s > today else 'expired' if e < today else 'current'].append(r)
    return out


is_active = lambda v: str(v or '').strip().lower() in ('active', 'yes', 'y', 'a', 'true', '1')   # S&D grids show Yes/No or Active/In-Active
found, prov = {}, {}


def put(key, value, src):
    if value not in (None, ''):
        found[key] = value
        prov[key] = {'source': src, 'resolved': now}


# ---- download template(s)
for x in sorted(glob.glob(os.path.join(disc, 'files', 'template*.xlsx'))):
    ws = openpyxl.load_workbook(x, data_only=True).active
    head = [str(c.value).replace('*', '').strip() if c.value is not None else '' for c in ws[1]]
    rows = [{h: ws.cell(row=r, column=i + 1).value for i, h in enumerate(head)} for r in range(2, ws.max_row + 1)
            if any(ws.cell(row=r, column=c).value not in (None, '') for c in range(1, len(head) + 1))]
    found['template_columns'] = head
    prov['template_columns'] = {'source': os.path.relpath(x, disc), 'resolved': now}
    found['template_rows'] = len(rows)
    for r in rows:
        put('entity_type', r.get('Entity Type'), os.path.basename(x))
        put('entity_code_in_template', r.get('Entity Code'), os.path.basename(x))
        put('template_subtype', r.get('Channel Hierarchy Code') or r.get('Channel Hierarchy'), os.path.basename(x))
        put('policy_with_uploads', r.get('Policy Code'), os.path.basename(x))
        if 'entity_type' in found:
            break

# ---- policies of the login distributor
pol = jload('policies')
if pol:
    c = classify(pol['rows'])
    is_active = lambda v: str(v or '').strip().lower() in ('active', 'yes', 'y', 'a', 'true', '1')      # S&D grids show Yes/No or Active/In-Active
    active = lambda rs: [r for r in rs if is_active(r.get('Status'))]
    put('current_policy', (active(c['current']) or c['current'] or [{}])[0].get('Policy code'), 'data/policies.json')
    put('future_policy', (active(c['future']) or c['future'] or [{}])[0].get('Policy code'), 'data/policies.json')
    put('expired_policy', (c['expired'] or [{}])[0].get('Policy code'), 'data/policies.json')
    found['policy_counts'] = {k: len(v) for k, v in c.items()}
    inact = [r for r in pol['rows'] if r.get('Status') and not is_active(r.get('Status'))]
    put('inactive_policy', (inact or [{}])[0].get('Policy code'), 'data/policies.json')

# ---- policies of another distributor (must not be visible to the login distributor)
po = jload('policies_other')
if po and pol:
    mine = {r.get('Policy code') for r in pol['rows']}
    theirs = [r.get('Policy code') for r in po['rows'] if r.get('Policy code') not in mine]
    put('other_dt_policy', theirs[0] if theirs else None, 'data/policies_other.json')

# ---- channel hierarchy (outlet subtypes)
ch = jload('channel_hierarchy')
if ch:
    is_active = lambda v: str(v or '').strip().lower() in ('active', 'yes', 'y', 'a', 'true', '1')
    act = [r for r in ch['rows'] if is_active(r.get('Status'))]
    ina = [r for r in ch['rows'] if r.get('Status') and not is_active(r.get('Status'))]
    used = {str(found.get('template_subtype') or '')}                      # already in the sample policy: keep for the "existing rows" cases
    pick = sorted((r for r in act if r.get('Code') not in used), key=lambda r: (not re.match(r'^[A-Z]\d+$', r.get('Code') or ''), r.get('Code') or ''))
    put('active_subtype', (pick[0] if pick else {}).get('Code'), 'data/channel_hierarchy.json')
    put('active_subtype_2', (pick[1] if len(pick) > 1 else {}).get('Code'), 'data/channel_hierarchy.json')
    put('inactive_subtype', (ina[0] if ina else {}).get('Code'), 'data/channel_hierarchy.json')
    found['channel_hierarchy_counts'] = {'active': len(act), 'inactive': len(ina), 'total': len(ch['rows'])}

json.dump({'run': os.path.basename(run), 'discovery': os.path.relpath(disc, os.path.dirname(run)), 'resolved': now,
           'values': found, 'provenance': prov}, open(os.path.join(run, 'data.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)

steps_p = os.path.join(run, 'steps.json')
filled, still = [], []
if os.path.exists(steps_p):
    doc = json.load(open(steps_p, encoding='utf-8'))
    d = doc.setdefault('defaults', {}).setdefault('data', {})
    for k, v in found.items():
        if k in d and d[k] in (None, '') and not isinstance(v, (dict, list)):
            filled.append(k)
            if apply:
                d[k] = v
    still = [k for k, v in d.items() if v in (None, '')]
    if apply:
        json.dump(doc, open(steps_p, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)

print('resolved:', {k: v for k, v in found.items() if not isinstance(v, list)})
print(('filled (applied): ' if apply else 'would fill: ') + (', '.join(filled) or 'nothing'))
print('still null in steps defaults: ' + (', '.join(still) or 'none'))
