"""Import the QA member's edited Excel back into the run.

  python runtime/qaos_import.py <file.xlsx> --run <run_dir> [--app snd] [--apply] [--approved-by "<who>"]

v4 approval workbook (sheets Questions / Cases / Steps / Data rows / Results, written by runtime/qaos_workbook.py):
  reads the QA member's edits - question answers and parked flags, case title / type / traces, path and AI case,
  step sentences and the event columns (event, element, data column, expect; inserted / deleted rows allowed), data rows
  (CASE_TYPE, EXPECTED_MESSAGE, data values) - re-runs the pre-run check (qaos_run.py check) on the edited documents and
  prints ONLY the diff and the new problems (decision 8.3). Without --apply nothing is written; with --apply (no
  problems) requirement.json / cases.json / steps.json / paths.json are updated (old copies in <run>/history/) and
  check.json is rewritten. Re-render the workbook afterwards with runtime/qaos_workbook.py.

Older step-draft workbook (sheet "Test Cases"): one row per case (TestCase#, Test Screen, Test Scenario, Description,
Pre-Requisite, Test Steps, Test Data, Expected Result, Priority). Each line of Test Steps is "<n>. [<Actor>] <step in the
standard wording>". Every step is checked with runtime/qaos_steps.py (verb, roles, screens, field/button labels against
the real screen). Without --apply nothing is written: it prints the problems (case, step number, reason, closest labels)
and the differences to the current draft. With --apply (only when there are no problems) it replaces
<run>/step_draft_<TC>.json (the old one goes to <run>/history/) and marks the approval as "imported from Excel"; pass
--approved-by only when the QA member has said in chat that the steps are final.
Exit code 0 = no problems.
"""
import argparse
import datetime
import json
import os
import re
import shutil
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qaos_steps as Q  # noqa: E402

LINE = re.compile(r'^\s*(\d+)\s*[.)]\s*(?:\[(?P<actor>[^\]]+)\]\s*)?(?P<step>.+?)\s*$')


def cell_lines(v):
    return [x for x in (str(v).splitlines() if v not in (None, '') else []) if x.strip()]


def read_cases(path):
    wb = load_workbook(path, data_only=True)
    if 'Test Cases' not in wb.sheetnames:
        sys.exit('sheet "Test Cases" not found (sheets: ' + ', '.join(wb.sheetnames) + ')')
    ws = wb['Test Cases']
    head = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
    need = ['TestCase#', 'Test Steps']
    for n in need:
        if n not in head:
            sys.exit(f'column "{n}" not found in the Test Cases sheet (columns: {", ".join(head)})')
    cases = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        g = lambda k: r[head[k]] if k in head and head[k] < len(r) else None
        if not g('TestCase#'):
            continue
        cases.append({'id': str(g('TestCase#')).strip(), 'screen': g('Test Screen'), 'scenario': g('Test Scenario'), 'title': g('Description'),
                      'preconditions': [x.lstrip('- ').strip() for x in cell_lines(g('Pre-Requisite'))], 'steps_text': cell_lines(g('Test Steps')),
                      'data': dict((k.strip(), v.strip()) for k, _, v in (x.partition(' = ') for x in cell_lines(g('Test Data'))) if v),
                      'expected': g('Expected Result'), 'priority': g('Priority')})
    return cases


def build(case, app):
    """-> (steps, problems). Steps use the same entry shape as qaos_steps.add()."""
    steps, problems, ctx, prev_actor = [], [], None, None
    for raw in case['steps_text']:
        m = LINE.match(raw)
        if not m:
            problems.append({'line': raw.strip(), 'reason': 'not in the form "<n>. [<Actor>] <step>"'})
            continue
        actor = (m.group('actor') or prev_actor or '').strip()
        text = m.group('step').strip()
        r = Q.check(text, app, ctx)
        if not r['ok']:
            problems.append({'n': int(m.group(1)), 'step': text, 'reason': '; '.join(r['problems']), 'try': r.get('try')})
            continue
        if r['verb'] in ('login', 'switch_user'):
            actor = r['slots']['role']
        ctx = r['screen'] or ctx
        prev_actor = actor
        steps.append({'n': len(steps) + 1, 'actor': actor, 'step': text, 'verb': r['verb'], 'slots': r['slots'], 'risk': r['risk'], 'screen': r['screen'] or ctx,
                      'expect': None, 'note': None, 'raw': raw.strip(), 'state': 'clear' if not r['warnings'] else 'unverified', 'warnings': r['warnings']})
    return steps, problems


def diff(old, new):
    o, n = [s['step'] for s in old], [s['step'] for s in new]
    out = []
    for i in range(max(len(o), len(n))):
        a = o[i] if i < len(o) else None
        b = n[i] if i < len(n) else None
        if a != b:
            out.append(f"step {i + 1}: {'(none)' if a is None else a}  ->  {'(none)' if b is None else b}")
    return out


# ---------------------------------------------------------------- v4: the ONE approval workbook (qaos_workbook.py)
# Decision 8.3: import the QA member's edits (cases, paths / AI case, step sentences and event columns, data rows,
# question answers), re-run the pre-run check and show only the diff and the new problems.
def _rows(ws):
    head = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        yield {h: (r[i] if i < len(r) else None) for i, h in enumerate(head) if h}


def _s(v):
    return '' if v is None else str(v).strip()


def _expect(v):
    """Steps sheet Expect cell -> step value: JSON object, '<type>: <text>' message, or plain text"""
    v = _s(v)
    if not v:
        return None
    if v.startswith('{'):
        try:
            return json.loads(v)
        except ValueError:
            pass
    m = re.match(r'^(toast|inline|popup|alert|message):\s*(.+)$', v, re.S)
    if m:
        return {'type': m.group(1) if m.group(1) != 'message' else 'toast', 'text': m.group(2).strip()}
    return v


def _num(v):
    """Excel turns codes like 15108843 into numbers: back to text without '.0'"""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return v if v is None or isinstance(v, str) else str(v)


def import_approval(xlsx, run, app, apply):
    import copy
    import difflib
    from jsonschema import Draft202012Validator
    import qaos_workbook as W
    d = W.load_run(run)
    R = d['R']
    if not d['cases']:
        sys.exit(f'cases.json not found in {run}')
    wb = load_workbook(xlsx, data_only=True)
    cases, steps, paths, req = (copy.deepcopy(d[k]) for k in ('cases', 'steps', 'paths', 'requirement'))
    old_check = d['check'] or R.run_check(run, app)
    diff, problems, touched = [], [], set()
    by_id = {c['id']: c for c in R.all_cases(cases)}
    backlog = {c['id'] for s in cases['sections'] if W.is_backlog(s) for c in s['cases']}

    # ---- Questions: answer / parked
    if 'Questions' in wb.sheetnames and req:
        amb = {a['id']: a for a in req.get('ambiguities', [])}
        for r in _rows(wb['Questions']):
            a = amb.get(_s(r.get('Id')))
            if not a:
                continue
            ans, parked = _s(r.get('Answer')), _s(r.get('Parked')).upper() == 'Y'
            if ans != _s(a.get('answer')):
                diff.append(f"question {a['id']}: answer {a.get('answer')!r} -> {ans or None!r}")
                a['answer'] = ans or None
                touched.add('requirement')
            was = a.get('parked') if a.get('parked') is not None else not a.get('answer')
            if parked != bool(was) or (ans and parked):
                a['parked'] = parked
                if parked != bool(was):
                    diff.append(f"question {a['id']}: parked {'Y' if was else '-'} -> {'Y' if parked else '-'}")
                touched.add('requirement')

    # ---- Cases: title, type, traces, path, AI case
    sheet_path = {}
    if 'Cases' in wb.sheetnames:
        for r in _rows(wb['Cases']):
            cid = _s(r.get('Id'))
            if not re.match(r'^TC\d{2,3}$', cid):
                continue
            c = by_id.get(cid)
            if not c:
                problems.append(f'Cases: {cid} is not in cases.json (new cases are designed with the test-designer, not typed in)')
                continue
            why = _s(c.get('why') or c.get('note'))
            shown = _s(('ALT-' if c['kind'] == 'negative' else '') + c['title'] + (f' - {why}' if cid in backlog and why else ''))
            title = _s(r.get('Title'))
            if title and title != shown:
                new = re.sub(r'^ALT-', '', title)
                if cid in backlog and why and new.endswith(f' - {why}'):
                    new = new[:-len(f' - {why}')]
                diff.append(f'case {cid}: title {c["title"]!r} -> {new!r}')
                c['title'] = new
                touched.add('cases')
            kind = _s(r.get('Type'))
            if kind and kind != c['kind']:
                if kind not in ('positive', 'negative', 'boundary', 'regression', 'business-effect'):
                    problems.append(f'Cases: {cid} type {kind!r} is not one of positive / negative / boundary / regression / business-effect')
                else:
                    diff.append(f"case {cid}: type {c['kind']} -> {kind}")
                    c['kind'] = kind
                    touched.add('cases')
            tr = [x.strip() for x in _s(r.get('Traces')).split(',') if x.strip()]
            if tr and tr != c['traces']:
                diff.append(f"case {cid}: traces {', '.join(c['traces'])} -> {', '.join(tr)}")
                c['traces'] = tr
                touched.add('cases')
            sheet_path[cid] = (_s(r.get('Path')) or None, _s(r.get('AI case')).upper() == 'Y')
    plist = R.path_list(paths) if paths else []
    if sheet_path and (paths is not None or any(p for p, _ in sheet_path.values())):
        if paths is None:
            paths = plist = []
        old = R.path_of(paths)
        for cid, (pid, ai) in sheet_path.items():
            opid = old.get(cid, (None, False, None))[0]
            if pid == opid:
                continue
            entry = cid
            for p in plist:                                  # take the case (and its row settings) out of its old path
                for k in ('rows', 'manual'):
                    for x in list(p.get(k) or []):
                        if R.row_case(x) == cid:
                            entry = x if k == 'rows' else cid
                            p[k].remove(x)
                if p.get('record') == cid:
                    p['record'] = None
            if pid:
                tgt = next((p for p in plist if p.get('path') == pid), None)
                if not tgt:
                    tgt = {'path': pid, 'label': '', 'record': None, 'rows': [], 'manual': []}
                    plist.append(tgt)
                tgt.setdefault('rows', []).append(entry)
            diff.append(f"case {cid}: path {opid or '-'} -> {pid or '-'}")
            touched.add('paths')
        for p in plist:                                      # AI case per path = the rows marked Y
            ys = [cid for cid, (pid, ai) in sheet_path.items() if pid == p.get('path') and ai]
            if ys == [p.get('record')] or (not ys and not any(c in sheet_path for c in [p.get('record')])):
                continue
            new = ys[0] if len(ys) == 1 else (ys or None)
            if new != p.get('record'):
                diff.append(f"path {p.get('path')}: AI case {p.get('record') or '-'} -> {new or '-'}")
                p['record'] = new
                touched.add('paths')

    # ---- Steps: sentences and event columns (aligned with difflib, so inserted / deleted rows do not shift the rest)
    if 'Steps' in wb.sheetnames and steps:
        old_rows = {}
        for row in W.step_rows(steps, d['check']):          # as the workbook rendered them (with or without check.json)
            old_rows.setdefault(row[0], []).append(row[2:])
        new_rows = {}
        for r in _rows(wb['Steps']):
            cid = _s(r.get('Case'))
            if not cid:
                continue
            new_rows.setdefault(cid, []).append((_s(r.get('Sentence')), _s(r.get('Event')), _s(r.get('Element')),
                                                 _s(r.get('Data column')), _s(r.get('Expect'))))
        sc = {c['case']: c for c in steps['cases']}
        for cid, rows in new_rows.items():
            if cid not in sc:
                problems.append(f'Steps: case {cid} has no steps entry in steps.json (add it to cases.json / steps first)')
                continue
            olds = [tuple(_s(x) for x in o) for o in old_rows.get(cid, [])]
            if rows == olds:
                continue
            ostep, out = sc[cid]['steps'], []
            for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=olds, b=rows, autojunk=False).get_opcodes():
                if tag == 'equal':
                    out += ostep[i1:i2]
                    continue
                for k in range(j2 - j1):
                    sen, ev, el, col, exp = rows[j1 + k]
                    base = copy.deepcopy(ostep[i1 + k]) if tag == 'replace' and i1 + k < i2 else {}
                    was = olds[i1 + k] if tag == 'replace' and i1 + k < i2 else None
                    if ev and ev not in R.EVENTS:
                        problems.append(f'Steps: {cid} step {len(out) + 1}: event {ev!r} is not one of {", ".join(R.EVENTS)}')
                    if base.get('do') and was and sen != was[0]:          # sentence rewritten: the QA wording wins over the v1 step
                        base = {k2: v for k2, v in base.items() if k2 in ('note', 'screen')}
                        base['event'] = ev or R.DO_EVENT.get(ostep[i1 + k].get('do'))
                    step = dict(base)
                    if sen:
                        step['sentence'] = sen
                    for key_, v in (('event', ev), ('element', el), ('data_column', col)):
                        if v or key_ in step:
                            step[key_] = v or None
                    if exp or 'expect' in step:
                        if not was or exp != was[4]:
                            step['expect'] = _expect(exp)
                    if not step.get('event'):
                        step.pop('event', None)
                    out.append(step)
                    diff.append(f"{cid} step {len(out)}: {(was[0] if was else '(new)')[:70]!r} -> {sen[:70]!r}"
                                + (f" [{' '.join(x for x in (ev, el, col) if x)}]" if (ev or el or col) else ''))
                if tag in ('delete', 'replace'):
                    for k in range(i1 + (j2 - j1), i2):
                        diff.append(f'{cid}: step removed: {olds[k][0][:70]!r}')
            sc[cid]['steps'] = out
            touched.add('steps')

    # ---- Data rows: CASE_TYPE / EXPECTED_MESSAGE (path row) and data values (case data row)
    if 'Data rows' in wb.sheetnames and steps:
        shown = {(p, c): (ct, msg, data) for p, ai, c, kind, ct, msg, data in W.data_rows(d)}
        sc = {c['case']: c for c in steps['cases']}
        fixed = set(W.DATA_FIXED)
        for r in _rows(wb['Data rows']):
            pid, cid = _s(r.get('Path')) or None, _s(r.get('Case'))
            if not cid or (pid, cid) not in shown:
                continue
            ct0, msg0, data0 = shown[(pid, cid)]
            ct, msg = _s(r.get('CASE_TYPE')), _s(r.get('EXPECTED_MESSAGE'))
            if (ct and ct != ct0) or msg != _s(msg0):
                for p in plist:
                    if p.get('path') != pid:
                        continue
                    for i, x in enumerate(p.get('rows') or []):
                        if R.row_case(x) == cid:
                            x = x if isinstance(x, dict) else {'case': cid}
                            if ct and ct != ct0:
                                x['CASE_TYPE'] = ct
                            if msg != _s(msg0):
                                x['EXPECTED_MESSAGE'] = msg
                            p['rows'][i] = x
                diff.append(f'{pid} {cid}: CASE_TYPE {ct0} -> {ct or ct0}, EXPECTED_MESSAGE {msg0!r} -> {msg!r}'
                            if ct and ct != ct0 else f'{pid} {cid}: EXPECTED_MESSAGE {msg0!r} -> {msg!r}')
                touched.add('paths')
            for col, v in r.items():
                if col in fixed or cid not in sc:
                    continue
                v = _num(v)
                if _s(v) != _s(data0.get(col)):
                    sc[cid].setdefault('data', {})[col] = v if _s(v) else None
                    diff.append(f'{cid} data {col}: {data0.get(col)!r} -> {v!r}')
                    touched.add('steps')

    # ---- schema checks of the edited documents, then the pre-run check on them
    schemas = os.path.join(R.SCHEMAS)
    for name, doc in (('cases', cases), ('steps', steps)):
        if name in touched:
            for e in list(Draft202012Validator(R.load(os.path.join(schemas, f'{name}.schema.json'))).iter_errors(doc))[:10]:
                problems.append(f'{name}.json schema: {list(e.absolute_path)}: {e.message[:160]}')
    new_check = R.run_check(run, app, cases=cases, steps=steps, paths=paths)
    problems += [f'paths: {e}' for e in new_check.get('paths_errors') or [] if e not in (old_check.get('paths_errors') or [])]

    print(f'import {os.path.basename(xlsx)} -> {run}')
    print(f'  {len(diff)} change(s)' + ('' if diff else ' - the workbook matches the run'))
    for x in diff[:200]:
        print('  CHANGED', x)
    newp = []
    for cid, r in new_check['cases'].items():
        o = (old_check.get('cases') or {}).get(cid) or {'readiness': None, 'reasons': []}
        added = [x for x in r['reasons'] if x not in o['reasons']]
        if r['readiness'] != o['readiness'] or added:
            newp.append(f"{cid}: {o['readiness'] or '(new)'} -> {r['readiness']}" + (f": {'; '.join(added[:3])}" if added else ''))
    for x in problems:
        print('  PROBLEM', x)
    for x in newp:
        print('  CHECK  ', x)
    if not problems and not newp:
        print('  no new problems')
    if problems:
        print('  nothing written: fix the problems in the workbook and import again')
        return 1
    if not apply:
        print('  (dry run: nothing written; add --apply to write the changes)')
        return 0
    stamp = datetime.datetime.now().strftime('%Y%m%d-%H%M%S')
    os.makedirs(os.path.join(run, 'history'), exist_ok=True)
    for name, doc in (('requirement', req), ('cases', cases), ('steps', steps), ('paths', paths)):
        if name not in touched:
            continue
        p = os.path.join(run, f'{name}.json')
        if os.path.exists(p):
            shutil.copy(p, os.path.join(run, 'history', f'{name}_{stamp}.json'))
        R.save(p, doc)
        print(f'  written {p}')
    R.save(os.path.join(run, 'check.json'), new_check)
    print(f"  written {os.path.join(run, 'check.json')}  |  " + ', '.join(f'{k} {v}' for k, v in sorted(new_check['summary'].items())))
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--run', required=True)
    ap.add_argument('--app', default='snd')
    ap.add_argument('--apply', action='store_true')
    ap.add_argument('--approved-by')
    a = ap.parse_args()
    if {'Cases', 'Steps'} <= set(load_workbook(a.xlsx, read_only=True).sheetnames):   # v4 approval workbook
        sys.exit(import_approval(a.xlsx, a.run, a.app, a.apply))
    bad = 0
    for case in read_cases(a.xlsx):
        steps, problems = build(case, a.app)
        p = os.path.join(a.run, f"step_draft_{case['id']}.json")
        old = json.load(open(p, encoding='utf-8')) if os.path.exists(p) else {'steps': [], 'case': {}}
        ch = diff(old['steps'], steps) if not problems else []      # step numbers shift when a line is invalid: no diff until fixed
        print(f"\n== {case['id']}: {len(case['steps_text'])} line(s) read, {len(steps)} valid step(s), {len(problems)} problem(s)"
              + ('' if problems else f", {len(ch)} change(s) vs the current draft"))
        for x in problems:
            print('  PROBLEM', (f"step {x['n']}: " if 'n' in x else '') + (x.get('step') or x['line']) + '\n          -> ' + x['reason'] + (f"  (try: {x['try']})" if x.get('try') else ''))
        if problems:
            print('  (the list of changes is shown once the problems are fixed)')
        for c in ch[:25]:
            print('  CHANGED', c)
        for st in steps:
            for w in st['warnings']:
                print(f"  CHECK  step {st['n']}: {w}")
        if problems:
            bad += 1
            continue
        lint_path = os.path.join(a.run, f".import_{case['id']}.json")
        json.dump({'steps': steps}, open(lint_path, 'w', encoding='utf-8'))
        lint = Q.lint(lint_path, a.app)
        os.remove(lint_path)
        for w in lint['warnings']:
            print('  warning', w)
        if lint['problems']:
            bad += 1
            for x in lint['problems']:
                print('  PROBLEM', x)
            continue
        if not a.apply:
            print('  (dry run: nothing written; add --apply to replace the draft)')
            continue
        os.makedirs(os.path.join(a.run, 'history'), exist_ok=True)
        if os.path.exists(p):
            shutil.copy(p, os.path.join(a.run, 'history', f"step_draft_{case['id']}_{datetime.datetime.now():%Y%m%d-%H%M%S}.json"))
        oc = old.get('case', {})
        new_case = {**oc, 'id': case['id'], 'screen': case['screen'] or oc.get('screen'), 'scenario': case['scenario'] or oc.get('scenario'),
                    'title': case['title'] or oc.get('title'), 'preconditions': case['preconditions'] or oc.get('preconditions', []),
                    'expected': case['expected'] or oc.get('expected'), 'priority': case['priority'] or oc.get('priority'), 'data': case['data'] or oc.get('data', {})}
        if ch and len(steps) != len(old['steps']):
            new_case['flow_map_stale'] = True          # the step ranges of the framework mapping no longer match
        approval = ({'gate': 'G3', 'status': 'approved (Excel import)', 'by': a.approved_by, 'date': datetime.date.today().isoformat(),
                     'note': 'the QA member edited the Test Steps column and confirmed the steps are final'} if a.approved_by else
                    {'gate': 'G3', 'status': 'imported from Excel, not yet confirmed as final', 'by': '', 'date': datetime.date.today().isoformat()})
        json.dump({'case': new_case, 'steps': steps, 'approval': approval}, open(p, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
        print(f"  written {p}  |  approval: {approval['status']}" + ('  |  framework mapping is now STALE (step count changed)' if new_case.get('flow_map_stale') else ''))
    sys.exit(1 if bad else 0)


if __name__ == '__main__':
    main()
