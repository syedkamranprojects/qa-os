"""QA OS deterministic player — executes a step-DSL flow against a live app with Selenium. No LLM at runtime.

It is the only component that types credentials. They come from env vars (or the "env" block of the
workspace .claude/settings.local.json) and are never logged or written to run files.

Usage:
  python runtime/qaos_player.py apps/snd/flows/P0_open_sku_substitution_policy.json            # run headed
  python runtime/qaos_player.py <flow.json> --dry-run      # validate + resolve every reference, no browser
  python runtime/qaos_player.py <flow.json> --headless

Output: runs/<case>/<yyyymmdd-hhmmss>/result.json, log.txt, step screenshots.
"""
import argparse, collections, datetime as dt, json, os, re, sys, time, traceback

import yaml
from jsonschema import Draft202012Validator

sys.stdout.reconfigure(encoding='utf-8')
QA_OS = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
SCHEMA = os.path.join(QA_OS, 'plugins', 'qa-os', 'schemas', 'steps.schema.json')
SETTINGS_LOCAL = os.path.normpath(os.path.join(QA_OS, '..', '.claude', 'settings.local.json'))


class StepFailed(Exception):
    pass


class NeedsData(StepFailed):
    """A {{data.x}} placeholder has no value yet (data-engineer, P2)."""


class NeedsVerb(StepFailed):
    """The step uses a DSL verb the player does not implement yet."""


PLAYER_VERBS = {'login', 'logout', 'harvest_menu', 'harvest_i18n', 'navigate', 'enter', 'choose', 'set_date', 'click',
                'select_row', 'verify_ui', 'capture', 'screenshot', 'wait',
                'download', 'upload', 'edit_excel', 'harvest_grid', 'capture_message', 'harvest_screen'}


# ---------------------------------------------------------------- app pack
class AppPack:
    def __init__(self, app_id):
        self.dir = os.path.join(QA_OS, 'apps', app_id)
        self.cfg = yaml.safe_load(open(os.path.join(self.dir, 'app.yaml'), encoding='utf-8'))
        self.screens = {}
        sdir = os.path.join(self.dir, 'screens')
        for fn in os.listdir(sdir) if os.path.isdir(sdir) else []:
            if fn.endswith('.json'):
                s = json.load(open(os.path.join(sdir, fn), encoding='utf-8'))
                self.screens[s['option_id']] = s

    def db_screen(self, option_id):
        """DB-declared dynamic screen (knowledge/screens_db); its field locators are NOT verified live."""
        p = os.path.join(self.dir, 'knowledge', 'screens_db', f'{option_id}.json')
        return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else None

    def env(self, name):
        name = name or self.cfg['default_env']
        return name, self.cfg['environments'][name]

    def menu_file(self, env, user):
        return os.path.join(self.dir, 'knowledge', 'env', env, f'menu.{user}.json')

    def resolve_menu(self, path_text, env, user):
        """'Transaction > SKU Substitution Policy' -> (top_id, option_id, search_text, route)."""
        parts = [p.strip() for p in path_text.split('>')]
        for s in self.screens.values():
            if [p.lower() for p in s['menu']['path']] == [p.lower() for p in parts]:
                m = s['menu']
                return m['top_id'], m['option_id'], m['search_text'], s.get('route')
        mf = self.menu_file(env, user)
        if os.path.exists(mf):
            for e in json.load(open(mf, encoding='utf-8'))['entries']:
                if [p.lower() for p in e['path']] == [p.lower() for p in parts]:
                    return e['top_id'], e['id'], e['title'], e.get('route')
        raise StepFailed(f"menu path not found in screens or harvested menu: {path_text}")


def credentials(app, user=None):
    """(username, password). For a named user: env SD_PASSWORD_<USER> (e.g. SD_PASSWORD_KPO_MP), else the default
    SD_TEST_USER/SD_TEST_PASSWORD pair if that pair is for the same user. Values are never printed."""
    ue, pe = app.cfg['credentials']['user_env'], app.cfg['credentials']['password_env']
    prefix = app.cfg['credentials'].get('password_env_prefix', 'SD_PASSWORD_')
    env = dict(os.environ)
    if os.path.exists(SETTINGS_LOCAL):
        env = {**json.load(open(SETTINGS_LOCAL, encoding='utf-8')).get('env', {}), **env}
    if user:                                                # a role name (Maker, Checker) maps to its user via app.yaml
        for e in (app.cfg.get('environments') or {}).values():
            hit = next((([u] if isinstance(u, str) else list(u or []) or [None])[0]   # a role's first user is its default
                        for r, u in (e.get('roles') or {}).items() if r.lower() == user.lower()), None)
            if hit:
                user = hit
                break
    # Per-person file outside the repo: env QAOS_CREDENTIALS, else ~/.qa-os/credentials.json (see credentials.example.json).
    cred_file = os.environ.get('QAOS_CREDENTIALS') or os.path.join(os.path.expanduser('~'), '.qa-os', 'credentials.json')
    if os.path.exists(cred_file):
        cf = json.load(open(cred_file, encoding='utf-8'))
        entry = ((cf.get(app.cfg['id']) or {}).get('users') or {}).get(user) if user else None
        if entry and entry.get('password'):
            return user, entry['password']
    if user:
        pw = env.get(prefix + re.sub(r'\W', '_', user.upper()))
        if pw:
            return user, pw
        if (env.get(ue) or '').lower() == user.lower():
            return env[ue], env.get(pe)
        return user, None
    return env.get(ue), env.get(pe)


# ---------------------------------------------------------------- placeholders
def resolve(value, ctx):
    if isinstance(value, list):
        return [resolve(v, ctx) for v in value]
    if isinstance(value, dict):
        return {k: resolve(v, ctx) for k, v in value.items()}
    if not isinstance(value, str):
        return value

    def sub(m):
        expr = m.group(1).strip()
        dm = re.fullmatch(r'today\s*([+-]\s*\d+)?', expr)
        if dm:
            d = dt.date.today() + dt.timedelta(days=int((dm.group(1) or '0').replace(' ', '')))
            return d.strftime(ctx['date_format'])
        ns, _, key = expr.partition('.')
        if ns == 'captured' and ctx.get('dry'):
            return f'<captured.{key}>'
        src = {'data': ctx['data'], 'captured': ctx['captured']}.get(ns)
        if src is None or key not in src or src[key] in (None, ''):
            raise NeedsData(f'no value for {{{{{expr}}}}}')
        return str(src[key])
    return re.sub(r'\{\{(.+?)\}\}', sub, value)


# ---------------------------------------------------------------- browser layer
class Browser:
    def __init__(self, headless, run_dir, log):
        from selenium import webdriver
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        from selenium.webdriver.common.keys import Keys
        self.By, self.EC, self.Keys, self.WebDriverWait = By, EC, Keys, WebDriverWait
        opts = webdriver.ChromeOptions()
        if headless:
            opts.add_argument('--headless=new')
        opts.add_argument('--window-size=1600,900')
        self.dl_dir = os.path.join(run_dir, 'files')             # downloads and files to upload live here
        os.makedirs(self.dl_dir, exist_ok=True)
        opts.add_experimental_option('prefs', {'download.default_directory': self.dl_dir,
                                               'download.prompt_for_download': False,
                                               'download.directory_upgrade': True,
                                               'safebrowsing.enabled': True,
                                               'profile.default_content_setting_values.automatic_downloads': 1})
        self.d = webdriver.Chrome(options=opts)
        self.run_dir, self.log = run_dir, log

    def by(self, loc):
        (kind, val), = loc.items()
        return {'id': self.By.ID, 'css': self.By.CSS_SELECTOR, 'xpath': self.By.XPATH, 'name': self.By.NAME}[kind], val

    def wait(self, cond, t=20):
        return self.WebDriverWait(self.d, t).until(cond)

    def find(self, loc, t=20, clickable=False):
        c = self.EC.element_to_be_clickable if clickable else self.EC.presence_of_element_located
        return self.wait(c(self.by(loc)), t)

    def idle(self, busy_css, t=30):
        end = time.time() + t
        while time.time() < end:
            if not [e for e in self.d.find_elements(self.By.CSS_SELECTOR, busy_css) if e.is_displayed()]:
                return
            time.sleep(0.3)

    def click(self, loc, t=20):
        el = self.find(loc, t, clickable=True)
        try:
            el.click()
        except Exception:
            self.d.execute_script('arguments[0].click()', el)
        return el

    def type(self, loc, text, enter=False):
        el = self.find(loc, clickable=True)
        el.click()
        el.send_keys(self.Keys.CONTROL, 'a')
        el.send_keys(self.Keys.DELETE)
        el.send_keys(text)
        # DevExtreme commits a text box's value on "change" (blur); press Tab so the form model really holds the text
        el.send_keys(self.Keys.ENTER if enter else self.Keys.TAB)

    def shot(self, name):
        p = os.path.join(self.run_dir, f'{name}.png')
        self.d.save_screenshot(p)
        return os.path.basename(p)

    def quit(self):
        try:
            self.d.quit()
        except Exception:
            pass


def is_readonly(driver, el):
    return driver.execute_script(
        "const i=arguments[0];const h=i.closest('.dx-widget');"
        "return i.readOnly||i.disabled||!!(h&&/dx-state-(readonly|disabled)/.test(h.className));", el)


def is_enabled_button(driver, el):
    return driver.execute_script(
        "const b=arguments[0];const h=b.closest('.dx-button')||b;"
        "return !(b.disabled||b.getAttribute('aria-disabled')==='true'||/dx-state-disabled/.test(h.className)||/\\bdisabled\\b/.test(b.className));", el)


# ---------------------------------------------------------------- player
class Player:
    def __init__(self, flow, app, args, run_dir, log):
        self.flow, self.app, self.args, self.run_dir, self.log = flow, app, args, run_dir, log
        self.env_name, self.env = app.env(flow.get('env'))
        self.user = flow.get('user') or next(iter(self.env['users']))
        self.ctx = {'data': flow.get('data', {}), 'captured': {},
                    'date_format': app.cfg['widgets']['date']['format'].replace('%Y', '%Y')}
        self.screen = None
        self.b = None

    # -- helpers
    def field(self, label):
        if not self.screen or label not in self.screen['fields']:
            raise StepFailed(f"field '{label}' not known on current screen {self.screen and self.screen['option_id']}")
        f = self.screen['fields'][label]
        if 'locator' not in f:
            raise StepFailed(f"field '{label}' on {self.screen['option_id']} is DB-declared only; learn the screen live first")
        return f

    def button(self, label):
        if not self.screen or label not in self.screen.get('buttons', {}):
            raise StepFailed(f"button '{label}' not known on current screen {self.screen and self.screen['option_id']}")
        return self.screen['buttons'][label]

    def busy(self):
        self.b.idle(self.app.cfg['widgets']['busy']['css'])

    # -- verbs
    def do_login(self, st):
        lg, url = self.app.cfg['login'], self.env['url']
        user_name, pwd = credentials(self.app, self.user)
        if not pwd:
            raise StepFailed(f'no password for user {self.user}: set SD_PASSWORD_{re.sub(chr(92)+"W", "_", self.user.upper())} (or SD_TEST_USER/SD_TEST_PASSWORD for the default user) in env or settings.local.json')
        self.b.d.get(url)
        self.b.type(lg['username'], user_name)
        self.b.find(lg['password']).send_keys(pwd)          # never logged
        self.b.click(lg['submit'])
        dist = resolve(st.get('distributor') or self.env['users'].get(self.user, {}).get('distributor', ''), self.ctx)
        if dist:
            dl = lg['distributor']
            if dist.startswith('*'):                          # "*" = first available distributor, "*3" = the third; the full list is recorded
                idx = int(dist[1:] or 1) - 1
                box = self.b.find(dl['input'], t=60, clickable=True)
                box.click()
                self.b.wait(lambda d: d.find_elements(self.b.By.XPATH, "//div[@role='option']"), 15)
                names = self.b.d.execute_script("return [...document.querySelectorAll('div[role=option]')].map(e=>e.textContent.trim())")
                os.makedirs(os.path.join(self.run_dir, 'data'), exist_ok=True)
                json.dump({'user': self.user, 'env': self.env_name, 'options': names},
                          open(os.path.join(self.run_dir, 'data', 'distributors.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
                if idx >= len(names):
                    raise StepFailed(f'only {len(names)} distributor(s) offered; wanted #{idx + 1}: {names}')
                dist = names[idx].split('-')[0].strip()
                self.log(f"  distributors offered ({len(names)}): {names[:12]}")
                self.b.click({'xpath': f"//div[@role='option' and starts-with(normalize-space(.),'{dist}-')]"})
            else:
                opt = {'xpath': dl['option']['xpath'].replace('{code}', dist)}
                box = self.b.find(dl['input'], t=60, clickable=True)
                box.click()
                try:                                             # searchable select: typing filters the list
                    box.send_keys(dist)
                    self.b.find(opt, t=8, clickable=True)
                except Exception:                                # not searchable or no match yet: open via the arrow
                    self.b.click(dl['button'])
                self.b.click(opt)
            self.b.wait(lambda d: dist in (d.find_element(*self.b.by(dl['input'])).get_attribute('value') or ''), 10)
            self.b.click(dl['proceed'])
        self.b.wait(lambda d: lg['landed_url_contains'] in d.current_url, 60)
        self.busy()
        self.logged_user = user_name
        return f'logged in as {user_name}' + (f', distributor {dist}' if dist else '')

    def do_harvest_menu(self, st):
        api = self.app.cfg['menu']['api']
        rows = self.b.d.execute_async_script("""
            const done = arguments[arguments.length-1];
            const tok = (localStorage.getItem('accessToken')||'').replace(/^"|"$/g,'');
            fetch(arguments[0], {credentials:'include', headers: tok ? {Authorization:'Bearer '+tok} : {}})
              .then(r => r.json()).then(done).catch(e => done({error: String(e)}));""", api)
        if isinstance(rows, dict) and rows.get('error'):
            raise StepFailed(f'menu API failed: {rows["error"]}')
        by_id = {r['optionGroupId']: r for r in rows}

        def chain(r):
            out, seen = [], set()
            while r and r['optionGroupId'] not in seen:
                seen.add(r['optionGroupId'])
                out.insert(0, r)
                r = by_id.get(r.get('parentOptionGroupId'))
            return out
        entries = []
        for r in rows:
            c = chain(r)
            entries.append({'id': r['optionGroupId'], 'title': (r.get('longDescription') or '').strip(),
                            'path': [(x.get('longDescription') or '').strip() for x in c],
                            'top_id': c[0]['optionGroupId'], 'parent': r.get('parentOptionGroupId'),
                            'app_option': r.get('appOptionId'), 'route': r.get('pathAngular'),
                            'param': r.get('pageUrlParam'), 'level': r.get('level'), 'seq': r.get('seqNo')})
        mf = self.app.menu_file(self.env_name, self.user)
        os.makedirs(os.path.dirname(mf), exist_ok=True)
        json.dump({'env': self.env_name, 'user': self.user, 'harvested': dt.datetime.now().isoformat(timespec='seconds'),
                   'source': api, 'entries': sorted(entries, key=lambda e: e['path'])},
                  open(mf, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
        return f'{len(entries)} menu entries -> {os.path.relpath(mf, QA_OS)}'

    def do_harvest_i18n(self, st):
        lang = st.get('lang', 'en')
        data = self.b.d.execute_async_script("""
            const done = arguments[arguments.length-1];
            fetch(arguments[0], {credentials:'include'}).then(r => r.json()).then(done).catch(e => done({error: String(e)}));""",
            f"{self.env['url'].rstrip('/')}/asset/i18n/{lang}.json")
        if isinstance(data, dict) and 'error' in data and len(data) == 1:
            raise StepFailed(f'i18n fetch failed: {data["error"]}')
        path = os.path.join(self.app.dir, 'knowledge', 'env', self.env_name, f'i18n.{lang}.json')
        os.makedirs(os.path.dirname(path), exist_ok=True)
        json.dump(data, open(path, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)

        def count(o):
            return sum(count(v) if isinstance(v, dict) else 1 for v in o.values()) if isinstance(o, dict) else 0
        return f'{count(data)} label keys -> {os.path.relpath(path, QA_OS)}'

    def do_navigate(self, st):
        top_id, opt_id, search, route = self.app.resolve_menu(st['menu'], self.env_name, self.user)
        m = self.app.cfg['menu']
        self.b.click({'css': m['top_item']['css'].replace('{top_id}', top_id)})
        self.b.type(m['search'], search)
        self.b.click({'css': m['item']['css'].replace('{option_id}', opt_id)})
        if route:
            self.b.wait(lambda d: route in d.current_url, 30)
        self.busy()
        self.screen = self.app.screens.get(opt_id) or self.app.db_screen(opt_id)
        return f'opened {st["menu"]} ({opt_id})' + ('' if self.screen else ' — no screen knowledge yet')

    def do_enter(self, st):
        self.b.type(self.field(st['field'])['locator'], resolve(st['value'], self.ctx))
        return f"{st['field']} = {resolve(st['value'], self.ctx)}"

    def do_set_date(self, st):
        self.b.type(self.field(st['field'])['locator'], resolve(st['value'], self.ctx), enter=True)
        return f"{st['field']} = {resolve(st['value'], self.ctx)}"

    def do_choose(self, st):
        f, val = self.field(st['field']), resolve(st['value'], self.ctx)
        self.b.click(f['locator'])
        if val.startswith('*'):                              # "*" = first option offered, "*3" = the third; the choice is recorded
            idx = int(val[1:] or 1) - 1
            self.b.wait(lambda d: d.find_elements(self.b.By.CSS_SELECTOR, '.dx-list-item'), 15)
            time.sleep(0.4)
            names = self.b.d.execute_script("return [...document.querySelectorAll('.dx-list-item')].filter(e=>e.offsetParent).map(e=>e.textContent.trim())")
            if idx >= len(names):
                raise StepFailed(f"{st['field']} offers {len(names)} option(s), wanted #{idx + 1}: {names[:8]}")
            val = names[idx]
            self.ctx['captured'][st.get('as') or 'chosen_' + re.sub(r'\W+', '_', st['field'].lower())] = val
        self.b.click({'xpath': self.app.cfg['widgets']['select']['option'].replace('{value}', val)})
        self.busy()
        return f"{st['field']} = {val}"

    def do_click(self, st):
        self.b.click(self.button(st['button'])['locator'])
        self.busy()
        return f"clicked {st['button']}"

    def do_select_row(self, st):
        g = self.screen['grids'][st['grid']]
        self.b.click({'xpath': g['row_by_first_cell'].replace('{value}', resolve(st['value'], self.ctx))})
        self.busy()
        return f"selected row {resolve(st['value'], self.ctx)}"

    def do_capture(self, st):
        el = self.b.find(self.field(st['field'])['locator'])
        self.ctx['captured'][st['as']] = el.get_attribute('value')
        return f"{st['as']} = {self.ctx['captured'][st['as']]}"

    def do_verify_ui(self, st):
        e, d = resolve(st['expect'], self.ctx), self.b.d
        if 'url_contains' in e:
            if e['url_contains'] not in d.current_url:
                raise StepFailed(f"url {d.current_url} lacks {e['url_contains']}")
            return f"url contains {e['url_contains']}"
        if 'breadcrumb' in e:
            text = d.execute_script("const b=document.querySelector('.breadcrumb, [class*=breadcrumb]');return b?b.innerText:''")
            pos = 0
            for item in e['breadcrumb']:
                i = text.find(item, pos)
                if i < 0:
                    raise StepFailed(f"breadcrumb '{text.strip()}' lacks '{item}' in order")
                pos = i + len(item)
            return 'breadcrumb ' + ' > '.join(e['breadcrumb'])
        if 'result_file' in e:
            import openpyxl
            name = self.ctx['captured'].get('result_file')
            if not name:
                raise StepFailed('the app returned no file after the upload')
            ws = openpyxl.load_workbook(self.file_path(name), data_only=True).active
            text = ' '.join('' if c is None else str(c) for r in ws.iter_rows(values_only=True) for c in r).lower()
            need = e['result_file'].get('contains', []) if isinstance(e['result_file'], dict) else []
            miss = [w for w in need if str(w).lower() not in text]
            if miss:
                raise StepFailed(f"returned file {name} lacks {miss}; it says: {text[:300]}")
            return f"returned file {name} mentions {need}" if need else f"returned file {name} present"
        if 'excel' in e:
            import openpyxl
            ws = openpyxl.load_workbook(self.file_path(e['excel']), data_only=True).active
            head = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
            norm = lambda s: str(s).replace('*', '').strip().lower()      # "Entity Code*" == "Entity Code"
            actual = [norm(h) for h in head if h]
            if 'columns' in e:
                # each expected column must equal an actual one, or be a prefix of it ("Channel Hierarchy" ~ "Channel Hierarchy Code")
                missing = [c for c in e['columns'] if not any(a == norm(c) or a.startswith(norm(c) + ' ') for a in actual)]
                extra = [h for h in head if h and not any(norm(h) == norm(c) or norm(h).startswith(norm(c) + ' ') for c in e['columns'])]
                if missing or extra:
                    raise StepFailed(f"excel columns {head}: missing {missing}, unexpected {extra}")
            first = {norm(h): ws.cell(row=2, column=i + 1).value for i, h in enumerate(head) if h}
            for k, v in e.get('cells', {}).items():
                got = first.get(norm(k))
                if (str(got) if got is not None else '') != str(v):
                    raise StepFailed(f"excel cell '{k}' = {got!r}, expected {v!r}")
            if 'has_rows_for' in e:
                found = any(str(e['has_rows_for']) in str(c.value) for r in ws.iter_rows(min_row=2) for c in r)
                if not found:
                    raise StepFailed(f"no excel row mentions {e['has_rows_for']}")
            return f"excel {e['excel']} ok ({ws.max_row - 1} data rows; columns {head})"
        if 'grid' in e and 'row' in e:
            rows = d.execute_script("return [...document.querySelectorAll(arguments[0])].map(r => r.textContent)",
                                    self.app.cfg['widgets']['grid']['row'].replace("//tr[contains(@class,'dx-data-row')]", 'tr.dx-data-row'))
            found = any(e['row'] in r for r in rows)
            if found != e.get('present', True):
                raise StepFailed(f"row '{e['row']}' {'missing from' if not found else 'unexpectedly present in'} {e['grid']} grid ({len(rows)} rows)")
            return f"row {e['row']} {'present' if found else 'absent'} in {e['grid']}"
        if 'grid' in e:
            # textContent, not .text: columns scrolled out of a narrow grid have empty visible text
            hdr = [t for t in d.execute_script(
                "return [...document.querySelectorAll(arguments[0])].map(h => h.textContent.trim())",
                self.app.cfg['widgets']['grid']['header']) if t]
            missing = [c for c in e['columns'] if c not in hdr]
            if missing:
                raise StepFailed(f'grid columns missing {missing}; found {hdr}')
            return f"grid columns ok ({len(e['columns'])})"
        if 'fields_present' in e:
            for lab in e['fields_present']:
                self.b.find(self.field(lab)['locator'], t=10)
            return f"{len(e['fields_present'])} fields present"
        if 'field' in e:
            el = self.b.find(self.field(e['field'])['locator'])
            if e.get('state') == 'invalid':
                bad = d.execute_script("const i=arguments[0];const h=i.closest('.dx-widget');"
                                       "return i.getAttribute('aria-invalid')==='true'||!!(h&&/dx-invalid/.test(h.className));", el)
                if not bad:
                    raise StepFailed(f"{e['field']} expected to be flagged invalid, is not")
                return f"{e['field']} flagged invalid"
            ro = is_readonly(d, el)
            if 'state' in e and (e['state'] == 'readonly') != ro:
                raise StepFailed(f"{e['field']} expected {e['state']}, is {'readonly' if ro else 'editable'}")
            if 'value' in e and (el.get_attribute('value') or '') != e['value']:
                raise StepFailed(f"{e['field']} value {el.get_attribute('value')!r} != {e['value']!r}")
            return f"{e['field']} {e.get('state', '')} {e.get('value', '')}".strip()
        if 'button' in e:
            en = is_enabled_button(d, self.b.find(self.button(e['button'])['locator']))
            if (e['state'] == 'enabled') != en:
                raise StepFailed(f"button {e['button']} expected {e['state']}, is {'enabled' if en else 'disabled'}")
            return f"button {e['button']} {e['state']}"
        if 'message_kind' in e:
            # Learn-first assertion: a success or an error message appears; the exact text is recorded for QA to tighten.
            js = ("return [...document.querySelectorAll('.dx-toast-message,.toast-message,.toast,[role=alert],.swal2-html-container,.alert,.dx-popup-content')]"
                  ".filter(x=>x.offsetParent).map(x=>x.innerText.trim()).filter(Boolean).join(' | ')")
            end, text = time.time() + 20, ''
            while time.time() < end and not text:
                text = d.execute_script(js)
                time.sleep(0.4)
            if not text:
                raise StepFailed(f"expected a {e['message_kind']} message, none appeared within 20s")
            low = text.lower()
            err = any(w in low for w in ('error', 'invalid', 'not exist', 'not found', 'fail', 'mandatory', 'required', 'overlap',
                                         'already', 'cannot', "can't", 'must', 'not allowed', 'inactive', 'in-active', 'duplicate', 'wrong'))
            ok_ = any(w in low for w in ('success', 'saved', 'uploaded', 'updated', 'created', 'completed', 'done'))
            kind = 'error' if err else 'success' if ok_ else 'unknown'
            os.makedirs(os.path.join(self.run_dir, 'messages'), exist_ok=True)
            with open(os.path.join(self.run_dir, 'messages', f"step{self.i:02d}.txt"), 'w', encoding='utf-8') as f:
                f.write(text)
            if kind != e['message_kind']:
                raise StepFailed(f"expected a {e['message_kind']} message, got ({kind}): {text[:200]}")
            return f"{e['message_kind']} message: {text[:160]}"
        if 'message_contains' in e:
            self.b.wait(lambda dr: e['message_contains'].lower() in dr.execute_script(
                "return [...document.querySelectorAll('.dx-toast-message,.toast-message,[role=alert]')].map(x=>x.innerText).join(' ')").lower(), 15)
            return f"message contains {e['message_contains']!r}"
        raise StepFailed(f'unsupported expectation {e}')

    # -- files (download / edit / upload) -------------------------------------------------
    def file_path(self, name):
        """A file name in the run's files/ folder, or an absolute path."""
        return name if os.path.isabs(name) else os.path.join(self.b.dl_dir, name)

    def do_download(self, st):
        import glob
        before = set(glob.glob(os.path.join(self.b.dl_dir, '*')))
        self.b.click(self.button(st['button'])['locator'])
        end = time.time() + 45
        got = None
        while time.time() < end:
            new = [f for f in glob.glob(os.path.join(self.b.dl_dir, '*')) if f not in before and not f.endswith(('.crdownload', '.tmp'))]
            if new:
                got = max(new, key=os.path.getmtime)
                time.sleep(0.5)
                break
            time.sleep(0.5)
        if not got:
            raise StepFailed(f"no file arrived in {self.b.dl_dir} within 45s after clicking {st['button']}")
        target = self.file_path(st['as'])
        if os.path.abspath(got) != os.path.abspath(target):
            if os.path.exists(target):
                os.remove(target)
            os.replace(got, target)
        return f"downloaded {os.path.basename(target)} ({os.path.getsize(target)} bytes)"

    def do_edit_excel(self, st):
        import openpyxl
        path = self.file_path(resolve(st['file'], self.ctx))
        if 'headers' in st and not os.path.exists(path):                # build a workbook from scratch (e.g. a wrong template)
            wb0 = openpyxl.Workbook()
            wb0.active.append(list(st['headers']))
            wb0.save(path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        norm = lambda s: str(s).replace('*', '').strip().lower()          # headers may carry a required-marker: "Entity Code*"
        head = {norm(c.value): c.column for c in ws[1] if c.value not in (None, '')}
        rows = resolve(st['rows'], self.ctx)
        for extra in range(2, ws.max_row + 1):                     # existing data rows are replaced
            for c in range(1, ws.max_column + 1):
                ws.cell(row=extra, column=c).value = None
        for i, row in enumerate(rows, 2):
            for k, v in row.items():
                if norm(k) not in head:
                    raise StepFailed(f"column '{k}' not in {os.path.basename(path)}; has {list(head)}")
                ws.cell(row=i, column=head[norm(k)]).value = None if v == '' else v
        wb.save(path)
        return f"{os.path.basename(path)}: {len(rows)} data row(s) written"

    def do_upload(self, st):
        path = self.file_path(resolve(st['file'], self.ctx))
        if not os.path.exists(path):
            raise StepFailed(f'file to upload not found: {path}')
        By = self.b.By

        def file_input():
            for e in self.b.d.find_elements(By.CSS_SELECTOR, 'input[type=file]'):
                return e
        inp = file_input()
        if inp is None:                                            # some screens create the input on click
            self.b.click(self.button(st['button'])['locator'])
            self.b.wait(lambda d: file_input() is not None, 10)
            inp = file_input()
        self.b.d.execute_script("arguments[0].style.display='block';arguments[0].style.visibility='visible';", inp)
        import glob
        before = set(glob.glob(os.path.join(self.b.dl_dir, '*')))
        inp.send_keys(os.path.abspath(path))
        self.busy()
        got, end = None, time.time() + float(st.get('wait_download', 8))     # e.g. an error file the app downloads on rejection
        while time.time() < end and not got:
            new = [f for f in glob.glob(os.path.join(self.b.dl_dir, '*')) if f not in before and not f.endswith(('.crdownload', '.tmp'))]
            got = max(new, key=os.path.getmtime) if new else None
            time.sleep(0.4)
        extra = ''
        if got:
            time.sleep(0.5)
            self.ctx['captured']['result_file'] = os.path.basename(got)
            try:
                import openpyxl
                ws = openpyxl.load_workbook(got, data_only=True).active
                rows = [' | '.join('' if c is None else str(c) for c in r) for r in ws.iter_rows(values_only=True)]
                os.makedirs(os.path.join(self.run_dir, 'messages'), exist_ok=True)
                with open(os.path.join(self.run_dir, 'messages', 'result_file.txt'), 'w', encoding='utf-8') as f:
                    f.write(chr(10).join(rows))
                extra = f"; app returned {os.path.basename(got)}: " + ' // '.join(rows)[:400]
            except Exception as ex:
                extra = f"; app returned {os.path.basename(got)} (unreadable: {type(ex).__name__})"
        return f"uploaded {os.path.basename(path)}" + extra

    def do_capture_message(self, st):
        """Record the visible toast/alert text (learn success and error messages)."""
        js = ("return [...document.querySelectorAll('.dx-toast-message,.toast-message,.toast,[role=alert],.swal2-html-container,.alert')]"
              ".filter(e=>e.offsetParent).map(e=>e.innerText.trim()).filter(Boolean).join(' | ')")
        end, text = time.time() + float(st.get('timeout', 10)), ''
        while time.time() < end and not text:
            text = self.b.d.execute_script(js)
            time.sleep(0.4)
        if not text:
            # No toast: a save blocked by field validation shows its reason only when the invalid field is hovered.
            from selenium.webdriver import ActionChains
            reasons = []
            for el in self.b.d.find_elements(self.b.By.CSS_SELECTOR, '.dx-invalid'):
                try:
                    ActionChains(self.b.d).move_to_element(el).perform()
                    time.sleep(0.5)
                    msg = self.b.d.execute_script(
                        "return [...document.querySelectorAll('.dx-invalid-message-content,.dx-invalid-message')].filter(e=>e.offsetParent).map(e=>e.textContent.trim()).filter(Boolean)[0]||''")
                    label = self.b.d.execute_script("const w=arguments[0].closest('.dx-field,.row,.form-group,tr');return w?w.innerText.split('\\n')[0].trim():''", el)
                    reasons.append(f"{label or 'field'}: {msg or 'invalid (no message)'}")
                except Exception:
                    continue
            if reasons:
                text = 'VALIDATION: ' + ' | '.join(reasons)
        if not text:
            raise StepFailed('no message appeared')
        self.ctx['captured'][st['as']] = text
        os.makedirs(os.path.join(self.run_dir, 'messages'), exist_ok=True)
        with open(os.path.join(self.run_dir, 'messages', f"{st['as']}.txt"), 'w', encoding='utf-8') as f:
            f.write(text)
        return f"{st['as']} = {text[:200]}"

    def do_harvest_grid(self, st):
        """Save a DevExtreme grid's visible rows as JSON under the run's data/ folder (data-engineer input)."""
        page_js = """
            const hdr=[...document.querySelectorAll('.dx-datagrid-headers td[role=columnheader]')].map(h=>h.textContent.trim());
            const rows=[...document.querySelectorAll('tr.dx-data-row')].map(r=>[...r.querySelectorAll('td')].map(c=>c.textContent.trim()));
            return {headers:hdr, rows:rows};"""
        # Wait for the grid to actually render: this app's page spinner is not always the DevExtreme load panel.
        nodata_js = "return [...document.querySelectorAll('.dx-datagrid-nodata')].some(e=>e.offsetParent && e.textContent.trim())"
        end = time.time() + float(st.get('wait', 25))
        while time.time() < end:
            if self.b.d.execute_script("return document.querySelectorAll('tr.dx-data-row').length") or self.b.d.execute_script(nodata_js):
                break
            time.sleep(0.4)
        time.sleep(0.4)
        data = self.b.d.execute_script(page_js)
        if not data['rows'] and not self.b.d.execute_script(nodata_js):
            raise StepFailed('grid rendered no rows and no "no data" message within the wait time')
        seen = {json.dumps(r) for r in data['rows']}
        # -- virtual scrolling: scroll the rows view to the end, collecting rows as they render
        scroll_js = """
            const c=document.querySelector('.dx-datagrid-rowsview .dx-scrollable-container');
            if(!c) return {moved:false};
            const before=c.scrollTop; c.scrollTop=before+Math.max(200,c.clientHeight*0.8);
            return {moved:c.scrollTop>before, top:c.scrollTop, max:c.scrollHeight-c.clientHeight};"""
        idle = 0
        for _ in range(int(st.get('max_scrolls', 400))):
            info = self.b.d.execute_script(scroll_js)
            if not info.get('moved'):
                break
            time.sleep(0.35)
            self.busy()
            new = [r for r in self.b.d.execute_script(page_js)['rows'] if json.dumps(r) not in seen]
            seen.update(json.dumps(r) for r in new)
            data['rows'] += new
            idle = 0 if new else idle + 1
            if idle >= 3 and info['top'] >= info['max'] - 5:
                break
        self.b.d.execute_script("const c=document.querySelector('.dx-datagrid-rowsview .dx-scrollable-container'); if(c) c.scrollTop=0;")
        # -- paging: try (1) the largest page size, (2) the numbered page buttons, (3) any "next" control
        pager_js = """
            const p=document.querySelector('.dx-datagrid-pager, .dx-pager');
            if(!p) return {found:false};
            return {found:true, html:p.outerHTML.slice(0,1500),
                    sizes:[...p.querySelectorAll('.dx-page-size')].map(e=>e.textContent.trim()),
                    pages:[...p.querySelectorAll('.dx-page')].map(e=>e.textContent.trim())};"""
        pager = self.b.d.execute_script(pager_js)
        diag = {'pager': pager, 'rows_first_page': len(data['rows'])}
        if pager.get('found'):
            try:                                                   # (1) biggest page size ("All" or 100/200)
                self.b.d.execute_script("""
                    const s=[...document.querySelectorAll('.dx-pager .dx-page-size')];
                    const all=s.find(e=>/^all$/i.test(e.textContent.trim()))||s[s.length-1]; if(all) all.click();""")
                time.sleep(1.2)
                self.busy()
                more = self.b.d.execute_script(page_js)['rows']
                if len(more) > len(data['rows']):
                    data['rows'] = more
                    seen = {json.dumps(r) for r in more}
            except Exception as ex:
                diag['page_size_error'] = str(ex)[:200]
            # (2) numbered pages: this pager shows only "1 2 3 4 5 … 19" (no next arrow), so click Page 2, 3, … in order
            last = max([int(p) for p in pager.get('pages', []) if p.isdigit()] or [1])
            click_page = """
                const p=document.querySelector('.dx-pager');
                const n=p&&p.querySelector('.dx-page[aria-label="Page '+arguments[0]+'"]');
                if(!n) return false; n.click(); return true;"""
            click_next = """
                const p=document.querySelector('.dx-pager');
                const n=p&&(p.querySelector('.dx-next-button')||p.querySelector('[class*=next]'));
                if(!n||/dx-button-disable|dx-state-disabled|disabled/.test(n.className)) return false;
                n.click(); return true;"""
            n = 2
            while n <= min(last, int(st.get('max_pages', 300))):
                if not (self.b.d.execute_script(click_page, n) or self.b.d.execute_script(click_next)):   # (3) fall back to a next arrow
                    diag['stopped_at_page'] = n
                    break
                self.busy()
                time.sleep(0.5)
                new = [r for r in self.b.d.execute_script(page_js)['rows'] if json.dumps(r) not in seen]
                seen.update(json.dumps(r) for r in new)
                data['rows'] += new
                if n == last:                                          # the last page number may grow as the pager re-renders
                    last = max([last] + [int(x) for x in self.b.d.execute_script(pager_js).get('pages', []) if x.isdigit()])
                n += 1
            diag['pages_walked'] = n - 1
            self.b.d.execute_script(click_page, 1)
        diag['rows_total'] = len(data['rows'])
        cols = [h for h in data['headers'] if h]
        recs = [dict(zip(cols, r[:len(cols)])) for r in data['rows']]
        os.makedirs(os.path.join(self.run_dir, 'data'), exist_ok=True)
        out = os.path.join(self.run_dir, 'data', f"{st['as']}.json")
        json.dump({'grid': st.get('grid'), 'env': self.env_name, 'harvested': dt.datetime.now().isoformat(timespec='seconds'),
                   'columns': cols, 'rows': recs, 'diagnostics': diag}, open(out, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
        return f"{len(recs)} rows -> data/{st['as']}.json" + ('' if diag['pager'].get('found') else ' (no pager found)')

    def do_harvest_screen(self, st):
        """Read-only learning: record the visible fields, buttons, grids, tabs and (optionally) dropdown options of the current screen."""
        js = r"""
        const vis = e => !!(e.offsetParent || e.getClientRects().length);
        const txt = e => (e && (e.innerText || e.textContent) || '').replace(/\s+/g,' ').trim();
        const widgetOf = e => { const c=e.closest('.dx-widget'); const cl=(c&&c.className)||''; const m=cl.match(/dx-(selectbox|datebox|textbox|numberbox|checkbox|lookup|tagbox|autocomplete|textarea|switch|radiogroup|button)/); return m?m[1]:null; };
        const labelFor = e => {
          if (e.id) { const l=document.querySelector('label[for="'+CSS.escape(e.id)+'"]'); if (l) return txt(l); }
          let n=e.closest('.dx-field,.form-group,.row,tr,[class*=field],[class*=col]');
          for (let i=0;i<4 && n;i++,n=n.parentElement) {
            const l=[...n.querySelectorAll('label,.dx-field-label,.label,span.caption,td:first-child')].find(x=>vis(x)&&txt(x)&&!x.contains(e)&&txt(x).length<60);
            if (l) return txt(l);
          }
          return null; };
        const fields=[...document.querySelectorAll('input:not([type=hidden]),textarea,select')].filter(vis).map(e=>({
          label: labelFor(e), tag:e.tagName.toLowerCase(), type:e.type||null, id:e.id||null, name:e.name||null,
          widget: widgetOf(e), placeholder:e.placeholder||null,
          required: e.required || e.getAttribute('aria-required')==='true' || /\*/.test(labelFor(e)||''),
          readonly: e.readOnly || e.disabled || /dx-state-(readonly|disabled)/.test((e.closest('.dx-widget')||{}).className||''),
          value: e.type==='password' ? null : (e.value||null)
        }));
        const buttons=[...document.querySelectorAll('button,.dx-button,[role=button],a.btn')].filter(vis).map(b=>({
          text:txt(b), id:b.id||null, tag:b.tagName.toLowerCase(), title:b.getAttribute('title')||b.getAttribute('aria-label')||null,
          disabled: b.disabled || b.getAttribute('aria-disabled')==='true' || /dx-state-disabled|disabled/.test(b.className)
        })).filter(b=>b.text||b.title);
        const grids=[...document.querySelectorAll('.dx-datagrid')].filter(vis).map((g,i)=>({
          index:i, id:g.id||null,
          headers:[...g.querySelectorAll('.dx-datagrid-headers td[role=columnheader]')].map(h=>h.textContent.trim()).filter(Boolean),
          rows:g.querySelectorAll('tr.dx-data-row').length,
          first_row:[...(g.querySelector('tr.dx-data-row')||{querySelectorAll:()=>[]}).querySelectorAll('td')].map(c=>c.textContent.trim()).slice(0,12),
          pager: !!g.querySelector('.dx-pager')
        }));
        const tabs=[...document.querySelectorAll('[role=tab],.dx-tab,.nav-tabs a,.mat-tab-label')].filter(vis).map(t=>({text:txt(t),selected:/dx-tab-selected|active|selected/.test(t.className)}));
        const headings=[...document.querySelectorAll('h1,h2,h3,h4,.heading,.dx-form-group-caption,.card-title,.panel-title')].filter(vis).map(txt).filter(Boolean);
        const checks=[...document.querySelectorAll('.dx-checkbox,.dx-radiobutton,.dx-switch')].filter(vis).map(c=>({text:txt(c.closest('label,.dx-field,div')),checked:/dx-checkbox-checked|dx-radiobutton-checked|dx-switch-on-value/.test(c.className)}));
        return {url:location.href, title:document.title, breadcrumb:txt(document.querySelector('.breadcrumb,[class*=breadcrumb]')),
                headings, tabs, fields, buttons, grids, checks};"""
        info = self.b.d.execute_script(js)
        if st.get('open_selects'):                               # read-only: open each dropdown, read its options, close it
            By = self.b.By
            opts = {}
            for f in info['fields']:
                if f['widget'] not in ('selectbox', 'lookup', 'tagbox', 'autocomplete') or not f['id']:
                    continue
                try:
                    box = self.b.d.find_element(By.ID, f['id'])
                    box.click()
                    time.sleep(0.6)
                    names = self.b.d.execute_script(
                        "return [...document.querySelectorAll('.dx-list-item,.dx-item[role=option]')].filter(e=>e.offsetParent).map(e=>e.textContent.trim()).filter(Boolean)")
                    opts[f['id']] = names[: int(st.get('max_options', 40))]
                    box.send_keys(self.b.Keys.ESCAPE)
                    self.b.d.execute_script("document.body.click()")
                    time.sleep(0.3)
                except Exception as ex:
                    opts[f['id']] = f'unreadable: {type(ex).__name__}'
            info['options'] = opts
        os.makedirs(os.path.join(self.run_dir, 'data'), exist_ok=True)
        out = os.path.join(self.run_dir, 'data', f"screen_{st['as']}.json")
        info['env'], info['harvested'] = self.env_name, dt.datetime.now().isoformat(timespec='seconds')
        json.dump(info, open(out, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
        return (f"{len(info['fields'])} fields, {len(info['buttons'])} buttons, {len(info['grids'])} grids, "
                f"{len(info['tabs'])} tabs -> data/screen_{st['as']}.json")

    def do_screenshot(self, st):
        return 'saved ' + self.b.shot(f"step{self.i:02d}_screenshot")

    def do_wait(self, st):
        time.sleep(float(st['seconds']))
        return f"waited {st['seconds']}s"

    def do_logout(self, st):
        lo = self.app.cfg['login']['logout']
        who = getattr(self, 'logged_user', self.user).upper()
        self.b.click({'xpath': lo['open']['xpath'].replace('{user_upper}', who)})
        self.b.click(lo['click'])
        self.b.wait(lambda d: d.find_elements(*self.b.by(self.app.cfg['login']['username'])), 30)
        return 'logged out'

    # -- dry run: resolve everything that can be resolved without a browser
    def dry(self, st):
        v = st['do']
        if v not in PLAYER_VERBS:
            raise NeedsVerb(f"verb '{v}' not implemented in the player yet" + (f" ({st['note']})" if st.get('note') else ''))
        if v in ('download', 'upload'):
            self.button(st['button'])
            return f"{v} via {st['button']}: {st.get('as') or resolve(st['file'], self.ctx)}"
        if v == 'edit_excel':
            return f"{st['file']}: {len(resolve(st['rows'], self.ctx))} row(s)"
        if v in ('harvest_grid', 'capture_message', 'harvest_screen'):
            return f"{v} as {st['as']}"
        if v == 'select_row':
            if not self.screen or st['grid'] not in self.screen.get('grids', {}):
                raise StepFailed(f"grid '{st['grid']}' not known on current screen")
            return f"row {resolve(st['value'], self.ctx)} in {st['grid']}"
        if v == 'navigate':
            try:
                top, opt, search, route = self.app.resolve_menu(st['menu'], self.env_name, self.user)
            except StepFailed:
                if any(x['do'] == 'harvest_menu' for x in self.flow['steps'][:self.i - 1]):
                    return 'menu path resolves live, after harvest_menu'
                raise
            self.screen = self.app.screens.get(opt)
            return f'menu -> li#{top} / search "{search}" / li#{opt} / {route}'
        if v in ('enter', 'choose', 'set_date', 'capture'):
            return f"{st['field']} -> {self.field(st['field'])['locator']} value={resolve(st.get('value', ''), self.ctx)}"
        if v == 'click':
            return f"{st['button']} -> {self.button(st['button'])['locator']}"
        if v == 'verify_ui':
            e = resolve(st['expect'], self.ctx)
            for lab in e.get('fields_present', []) + ([e['field']] if 'field' in e else []):
                self.field(lab)
            if 'button' in e:
                self.button(e['button'])
            if 'grid' in e and (not self.screen or e['grid'] not in self.screen.get('grids', {})):
                raise StepFailed(f"grid '{e['grid']}' not known on current screen")
            return f'expect {e}'
        if v == 'login':
            u, p = credentials(self.app, self.user)
            return f"login recipe ok; credentials for {self.user} {'present' if u and p else 'MISSING'}"
        return 'ok'

    def run(self, dry_run=False):
        results, verdict = [], 'pass'
        self.ctx['dry'] = dry_run
        if not dry_run:
            self.b = Browser(self.args.headless, self.run_dir, self.log)
        try:
            for self.i, st in enumerate(self.flow['steps'], 1):
                t0 = time.time()
                try:
                    if not dry_run and st['do'] not in PLAYER_VERBS:
                        raise NeedsVerb(f"verb '{st['do']}' not implemented in the player yet")
                    msg = self.dry(st) if dry_run else getattr(self, 'do_' + st['do'])(st)
                    results.append({'n': self.i, 'do': st['do'], 'status': 'pass', 'detail': msg, 'secs': round(time.time() - t0, 1)})
                    self.log(f"  {self.i:02d} PASS {st['do']:<12} {msg}")
                except (NeedsData, NeedsVerb) as ex:
                    status = 'needs-data' if isinstance(ex, NeedsData) else 'needs-verb'
                    results.append({'n': self.i, 'do': st['do'], 'status': status, 'detail': str(ex)})
                    self.log(f"  {self.i:02d} {status.upper()} {st['do']:<12} {ex}")
                    if dry_run:                              # keep checking the remaining steps
                        verdict = verdict if verdict == 'fail' else status
                        continue
                    verdict = status
                    break
                except Exception as ex:
                    verdict = 'fail'
                    shot = None if dry_run else self.b.shot(f'step{self.i:02d}_FAIL')
                    detail = str(ex) if isinstance(ex, StepFailed) else f'{type(ex).__name__}: {ex}'
                    results.append({'n': self.i, 'do': st['do'], 'status': 'fail', 'detail': detail, 'evidence': shot,
                                    'secs': round(time.time() - t0, 1)})
                    self.log(f"  {self.i:02d} FAIL {st['do']:<12} {detail}")
                    if not isinstance(ex, StepFailed):
                        self.log(traceback.format_exc(limit=3))
                    break
            if verdict == 'fail' and not dry_run and self.flow['steps'][-1]['do'] == 'logout':
                try:
                    self.do_logout({})
                    self.log('  (logged out after failure)')
                except Exception:
                    pass
        finally:
            if self.b:
                self.b.quit()
        return verdict, results


def load_cases(doc):
    """A single-case file, or a bundle (steps.json) whose defaults are merged into each case."""
    if 'cases' not in doc:
        return None, [doc]
    out = []
    for c in doc['cases']:
        merged = {'app': doc['app'], 'env': doc.get('env'), 'user': doc.get('user'), **{k: v for k, v in c.items() if v is not None}}
        merged['data'] = {**doc.get('defaults', {}).get('data', {}), **c.get('data', {})}
        out.append(merged)
    return doc['key'], out


def run_case(flow, flow_path, key, args, stamp):
    app = AppPack(flow['app'])
    base = os.path.dirname(flow_path)
    if os.path.exists(os.path.join(base, 'run.json')):              # steps.json inside a QA OS run folder
        run_dir = os.path.join(base, 'exec', flow['case'], stamp + ('-dry' if args.dry_run else ''))
    else:
        run_dir = os.path.join(QA_OS, 'runs', *([key] if key else []), flow['case'], stamp + ('-dry' if args.dry_run else ''))
    os.makedirs(run_dir, exist_ok=True)
    log_f = open(os.path.join(run_dir, 'log.txt'), 'w', encoding='utf-8')

    def log(m):
        if not (args.dry_run and args.quiet):
            print(m)
        log_f.write(m + '\n')
        log_f.flush()

    log(f"{'DRY RUN' if args.dry_run else 'RUN'} {flow['case']} — {flow.get('title', '')}")
    p = Player(flow, app, args, run_dir, log)
    log(f"  app={flow['app']} env={p.env_name} url={p.env['url']} user={p.user}")
    started = dt.datetime.now().isoformat(timespec='seconds')
    verdict, results = p.run(dry_run=args.dry_run)
    json.dump({'key': key, 'case': flow['case'], 'title': flow.get('title'), 'app': flow['app'], 'env': p.env_name,
               'user': p.user, 'flow': os.path.relpath(flow_path, QA_OS), 'mode': 'dry-run' if args.dry_run else 'live',
               'readiness': flow.get('readiness'), 'started': started,
               'finished': dt.datetime.now().isoformat(timespec='seconds'), 'verdict': verdict, 'steps': results},
              open(os.path.join(run_dir, 'result.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
    log(f"VERDICT: {verdict.upper()}  ->  {os.path.relpath(run_dir, QA_OS)}")
    return verdict, run_dir, results


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('flow', help='single-case flow file or a steps.json bundle')
    ap.add_argument('--case', action='append', help='case id(s) to run from a bundle (repeatable); default: all')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--headless', action='store_true')
    ap.add_argument('--force', action='store_true', help='live-run cases even if readiness is not "ready"')
    ap.add_argument('--quiet', action='store_true', help='dry-run: print only the summary table')
    args = ap.parse_args()

    flow_path = os.path.abspath(args.flow)
    doc = json.load(open(flow_path, encoding='utf-8'))
    errs = sorted(Draft202012Validator(json.load(open(SCHEMA, encoding='utf-8'))).iter_errors(doc), key=str)
    if errs:
        sys.exit('flow does not match steps.schema.json:\n' + '\n'.join(
            f" - {list(e.absolute_path)}: {e.message[:300]}" for e in errs[:20]))

    key, cases = load_cases(doc)
    if args.case:
        cases = [c for c in cases if c['case'] in args.case]
        if not cases:
            sys.exit(f'no such case(s): {args.case}')
    stamp = dt.datetime.now().strftime('%Y%m%d-%H%M%S')
    summary = []
    for c in cases:
        if not args.dry_run and c.get('readiness', 'ready') != 'ready' and not args.force:
            print(f"SKIP {c['case']}: readiness={c.get('readiness')} blocked_by={c.get('blocked_by')}")
            summary.append((c['case'], 'skipped', c.get('readiness'), None))
            continue
        verdict, run_dir, results = run_case(c, flow_path, key, args, stamp)
        blockers = sorted({r['detail'] for r in results if r['status'] in ('needs-data', 'needs-verb', 'fail')})
        summary.append((c['case'], verdict, c.get('readiness'), blockers))
    if len(summary) > 1:
        print(f"\n{'case':<8} {'verdict':<11} {'declared':<15} blockers")
        for case, verdict, declared, blockers in summary:
            print(f"{case:<8} {verdict:<11} {str(declared):<15} {'; '.join(blockers or [])[:150]}")
        counts = collections.Counter(v for _, v, _, _ in summary)
        print('TOTAL', dict(counts))
    sys.exit(0 if all(v in ('pass', 'skipped') or (args.dry_run and v in ('needs-data', 'needs-verb'))
                      for _, v, _, _ in summary) else 1)


if __name__ == '__main__':
    main()
